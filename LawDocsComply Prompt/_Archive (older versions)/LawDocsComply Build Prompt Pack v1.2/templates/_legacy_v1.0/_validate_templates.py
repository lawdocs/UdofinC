# -*- coding: utf-8 -*-
"""Validate LawDocsComply question templates (LDC-TPL-*-v1.0) against the owner's base layout.

Usage: python3 _validate_templates.py <templates_dir> <owner_sample_1.xlsx> [<owner_sample_2.xlsx> ...] > _validation_report.md
Checks (each must PASS for every template):
  T01 sheets present and ordered (Sheet1, Rough, _Template_Info)
  T02 row 1 and row 2 of Sheet1 identical to every owner sample (values) and row-2 SHA-256 matches
  T03 merged ranges in row 1 identical to the owner sample
  T04 freeze panes at A3
  T05 Sheet1 has no data at or below row 3
  T06 seven family-prefixed named ranges resolve to contiguous, non-empty, duplicate-free lists on Rough, each preceded by its header row
  T07 base vocabulary included (frequencies, departments, compliance types, risk = High/Medium/Low)
  T08 list validations on C, D, H, M, O, P, R covering rows 3-2000 and pointing to the right named range
  T09 _Template_Info present, protected, template code consistent with file name and custom properties, header hash recorded correctly
  T10 IFSCA-GEN: Rough and the six owner names byte-identical to the owner's sample
  T11 row-2 header font is Times New Roman bold (style preserved)
"""
import glob
import hashlib
import os
import sys

import openpyxl
from openpyxl.utils import get_column_letter

COLS = [get_column_letter(i) for i in range(1, 31)]
SEP = '\x1f'
FREQ_BASE = ['Annually', 'Half Yearly', 'Quarterly', 'Monthly', 'Fortnightly', 'Weekly', 'Daily', 'Event Based', 'Immediately', 'Ongoing/Continuous', 'One-time']
DEPT_BASE = ['Governance & Administration', 'Finance, Treasury and Accounts', 'Legal & Litigation', 'IT', 'Operations & Management', 'HR', 'Compliance', 'Returns & Reporting', 'Risk Management']
COMP_BASE = ['Policies', 'Compliance', 'Returns & Reporting', 'Committee', 'Internal Compliance', 'Suggestive Compliance']
KEYS = {'Freq': 'H', 'Sector': 'C', 'Entity': 'D', 'Dept': 'M', 'CompType': 'O', 'Risk': 'R', 'Filing': 'P'}


def rowvals(ws, r):
    return ['' if ws[f'{c}{r}'].value is None else str(ws[f'{c}{r}'].value) for c in COLS]


def h(vals):
    return hashlib.sha256(SEP.join(vals).encode('utf-8')).hexdigest()


def strip(vals):
    vals = list(vals)
    while vals and vals[-1] is None:
        vals.pop()
    return vals


def resolve(wb, name):
    dn = wb.defined_names.get(name)
    if dn is None:
        return None
    sheet, ref = list(dn.destinations)[0]
    a, b = ref.replace('$', '').split(':')
    return sheet, int(a[1:]), int(b[1:])


def main():
    tdir, samples = sys.argv[1], sys.argv[2:]
    base = [openpyxl.load_workbook(s) for s in samples]
    b0 = base[0]['Sheet1']
    row1, row2 = rowvals(b0, 1), rowvals(b0, 2)
    for b in base[1:]:
        assert rowvals(b['Sheet1'], 1) == row1 and rowvals(b['Sheet1'], 2) == row2, 'owner samples disagree on rows 1-2'
    merged0 = sorted(str(m) for m in b0.merged_cells.ranges)
    canon = h(row2)
    owner_rough = [c.value for c in base[-1]['Rough']['A']]
    owner_names = {n: d.attr_text for n, d in base[-1].defined_names.items()}

    files = sorted(glob.glob(os.path.join(tdir, 'LDC-TPL-*.xlsx')))
    print('# Template validation report\n')
    print(f'Validator: `templates/_validate_templates.py` · Owner samples: {", ".join(os.path.basename(s) for s in samples)}  ')
    print(f'Canonical row-2 header SHA-256: `{canon}`  ')
    print(f'Templates checked: {len(files)}\n')
    print('| Template | T01 | T02 | T03 | T04 | T05 | T06 | T07 | T08 | T09 | T10 | T11 | Result |')
    print('|---|---|---|---|---|---|---|---|---|---|---|---|---|')
    fails = []
    details = []
    for f in files:
        name = os.path.basename(f)
        code = name.split(' - ')[0]
        fam = code[len('LDC-TPL-'):-len('-v1.0')]
        wb = openpyxl.load_workbook(f)
        res = {}
        msgs = []
        res['T01'] = wb.sheetnames == ['Sheet1', 'Rough', '_Template_Info']
        ws = wb['Sheet1']
        res['T02'] = rowvals(ws, 1) == row1 and rowvals(ws, 2) == row2 and h(rowvals(ws, 2)) == canon
        res['T03'] = sorted(str(m) for m in ws.merged_cells.ranges) == merged0
        res['T04'] = ws.freeze_panes == 'A3'
        res['T05'] = all(v is None for row in ws.iter_rows(min_row=3, values_only=True) for v in row)
        # names
        prefix = None
        for n in wb.defined_names.keys():
            if n.endswith('_Freq'):
                prefix = n[:-5]
        ok6 = prefix is not None
        lists = {}
        rough = wb['Rough']
        if ok6:
            for k in KEYS:
                rr = resolve(wb, f'{prefix}_{k}')
                if rr is None or rr[0] != 'Rough':
                    ok6 = False
                    msgs.append(f'name {prefix}_{k} missing')
                    continue
                _, a, b = rr
                vals = [rough[f'A{i}'].value for i in range(a, b + 1)]
                hdr = rough[f'A{a - 1}'].value
                if any(v is None or str(v).strip() == '' for v in vals) or len(set(vals)) != len(vals) or not hdr:
                    ok6 = False
                    msgs.append(f'{k}: blank/duplicate/missing header')
                if rough[f'A{b + 1}'].value is not None:
                    ok6 = False
                    msgs.append(f'{k}: list not terminated by blank row')
                lists[k] = vals
        res['T06'] = ok6
        res['T07'] = ok6 and set(FREQ_BASE) <= set(lists.get('Freq', [])) and DEPT_BASE == lists.get('Dept', [])[:9] and \
            COMP_BASE == lists.get('CompType', [])[:6] and lists.get('Risk') == ['High', 'Medium', 'Low']
        if not res['T07']:
            msgs.append('base vocabulary incomplete or reordered')
        # DV
        dvs = {}
        for dv in ws.data_validations.dataValidation:
            dvs[str(dv.sqref)] = (dv.type, dv.formula1)
        ok8 = True
        for k, col in KEYS.items():
            want = (f'{col}3:{col}2000')
            if want not in dvs or dvs[want] != ('list', f'{prefix}_{k}'):
                ok8 = False
                msgs.append(f'DV {want} missing/wrong: {dvs.get(want)}')
        res['T08'] = ok8
        info = wb['_Template_Info']
        infotext = {str(c.value) for row in info.iter_rows() for c in row if c.value is not None}
        props = {p.name: p.value for p in wb.custom_doc_props.props}
        res['T09'] = info.protection.sheet and code in infotext and canon in infotext and props.get('LDC_TemplateCode') == code and \
            props.get('LDC_HeaderSHA256') == canon and wb.properties.title == code
        if fam == 'IFSCA-GEN':
            names_ok = all(wb.defined_names.get(n) is not None and wb.defined_names[n].attr_text == t for n, t in owner_names.items())
            res['T10'] = strip([c.value for c in rough['A']]) == strip(owner_rough) and names_ok
        else:
            res['T10'] = None
        c2 = ws['G2']
        res['T11'] = c2.font.name == 'Times New Roman' and bool(c2.font.b)
        passed = all(v is not False for v in res.values())
        if not passed:
            fails.append(name)
        cells = ['PASS' if res[k] else ('n/a' if res[k] is None else 'FAIL') for k in sorted(res)]
        print(f'| {code} | ' + ' | '.join(cells) + f' | {"PASS" if passed else "FAIL"} |')
        details.append((code, prefix, {k: len(v) for k, v in lists.items()}, msgs))
    print('\n## List sizes per template\n')
    print('| Template | Name prefix | H freq | C | D | M dept | O type | R risk | P filing | Messages |')
    print('|---|---|---|---|---|---|---|---|---|---|')
    for code, prefix, n, msgs in details:
        print(f"| {code} | `{prefix}_*` | {n.get('Freq')} | {n.get('Sector')} | {n.get('Entity')} | {n.get('Dept')} | {n.get('CompType')} | {n.get('Risk')} | {n.get('Filing')} | {'; '.join(msgs) or '—'} |")
    print(f'\n**Overall: {"PASS" if not fails else "FAIL (" + ", ".join(fails) + ")"}** — {len(files) - len(fails)}/{len(files)} templates passed all checks.')
    return 1 if fails else 0


if __name__ == '__main__':
    sys.exit(main())
