import sys, re
sys.path.insert(0,'/home/claude/ldc_v11/catalogue/_build')
from rbi_data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

CAT = {  # category -> (abbr, P15 family, P15 type, template, command, wave)
 'Commercial Banks':('SCB','BANK','SCB_PSB / SCB_PVT / SCB_FOREIGN','LDC-TPL-RBI-SCB','01_RBI-SCB_Command.md','W3'),
 'Small Finance Banks':('SFB','BANK','SFB','LDC-TPL-RBI-SCB','01_RBI-SCB_Command.md','W3'),
 'Payments Banks':('PB','BANK','PB','LDC-TPL-RBI-SCB','01_RBI-SCB_Command.md','W3'),
 'Local Area Banks':('LAB','BANK','LAB','LDC-TPL-RBI-SCB','01_RBI-SCB_Command.md','W3'),
 'Regional Rural Banks':('RRB','BANK','RRB','LDC-TPL-RBI-SCB','01_RBI-SCB_Command.md','W3'),
 'Urban Co-operative Banks':('UCB','UCB','UCB_STATE / UCB_MSCS / UCB_SALARY_EARNERS / UCB_UNIT','LDC-TPL-RBI-UCB','02_RBI-UCB_Command.md','W2'),
 'Rural Co-operative Banks':('RCB','RCB','STCB / DCCB','LDC-TPL-RBI-RCB','03_RBI-RCB_Command.md','W2'),
 'All India Financial Institutions':('AIFI','AIFI','NABARD / SIDBI / EXIM / NHB / NaBFID','LDC-TPL-RBI-OTH','06_RBI-OTH_Command.md','W3'),
 'Non-Banking Financial Companies':('NBFC','NBFC','All NBFC types (see applicability)','LDC-TPL-RBI-NBFC','04_RBI-NBFC_Command.md','W1'),
 'Asset Reconstruction Companies':('ARC','ARC','ARC','LDC-TPL-RBI-ARC','05_RBI-ARC_Command.md','W2'),
 'Credit Information Companies':('CIC-CI','PAY (RBI-OTH: credit information company)','CREDIT_INFORMATION_COMPANY','LDC-TPL-RBI-OTH','06_RBI-OTH_Command.md','W2'),
 'Non-bank PPI Issuers':('PPI','PAY','PPI_NONBANK','LDC-TPL-RBI-OTH','06_RBI-OTH_Command.md','W3'),
}
NBFC_TYPE={'Housing Finance Companies':('HFC','HFC'),'Core Investment Companies':('CIC_CORE','CIC'),'Mortgage Guarantee Companies':('MGC','MGC'),
 'Standalone Primary Dealers':('SPD','SPD'),'Non-Operative Financial Holding Companies':('NOFHC','NOFHC'),'Microfinance Institution':('MFI','MFI'),
 'Peer to Peer Lending Platform':('P2P','P2P'),'Account Aggregator':('AA','AA')}
NBFC_APPL={ # secondary source: Vinod Kothari summary (Dec 2025) -> [VERIFY against each Direction's applicability para]
 'Registration, Exemptions and Framework for Scale Based Regulation':('All NBFCs incl. HFCs; layer determination','SBR Directions Ch I–III, XII (para 105), XVI, V (para 115); Master Circular on exemptions from RBI Act'),
 'Acquisition of Shareholding or Control':('All NBFCs incl. HFCs','SBR para 42; HFC Directions Ch VIII'),
 'Voluntary Amalgamation':('All NBFCs','RBI FAQs 84–86; SBR para 42 (relevant parts)'),
 'Branch Authorisation':('All NBFCs incl. HFCs','SBR paras 37.6, 43, 46'),
 'Undertaking of Financial Services':('NBFC-D, ICC, MFI, Factor, IDF-NBFC, IFC, HFCs; paras 7–12 from 1 Jan 2026','Investment in AIF Directions 2025; SBR para 32C, Annex XVI, XVIII, Sec. VIII; HFC Annex XVI'),
 'Governance':('All NBFCs except HFCs; BL Ch III; ML Ch II–IV; UL Ch II–V (para 3 — confirmed on RBI page)','SBR Ch XI (governance); repeal para 44'),
 'Prudential Norms on Capital Adequacy':('NBFC-ICC, Factor, deposit-taking, IFC, IDF (others per type directions)','SBR Ch IV para 9.2, Ch IX paras 81.1–86.3, Annex II para 3, Annex XIV; SSA Ch VI'),
 'Acceptance of Public Deposits':('Deposit-taking NBFCs (not HFCs, insurance cos., non-deposit-taking ICCs)','MD Acceptance of Public Deposits (Reserve Bank) Directions, 2016'),
 'Credit Facilities':('All NBFCs incl. HFCs','Digital Lending, Gold/Silver Collateral, Project Finance, Non-fund based credit facilities Directions, 2025'),
 'Credit Cards: Issuance and Conduct':('NBFCs permitted to issue credit cards','MD Credit Card and Debit Card – Issuance and Conduct, 2022'),
 'Credit Information Reporting':('All NBFCs incl. HFCs','MD Credit Information Reporting, 2025'),
 'Credit Risk Management':('All NBFCs incl. HFCs','SBR paras 93, 40, 35, 52, 65, 68; LEI circular 2017; CERSAI filing'),
 'Concentration Risk Management':('All NBFCs','SBR paras 32 (except 32C), 34, 91, 110'),
 'Transfer and Distribution of Credit Risk':('All NBFCs; Part B from 1 Jan 2026','MD Transfer of Loan Exposures 2021; Co-Lending Directions 2025'),
 'Securitisation Transactions':('All NBFCs incl. HFCs','MD Securitisation of Standard Assets (except capital computation and disclosure)'),
 'Classification, Valuation and Operation of Investment Portfolio':('All NBFCs incl. HFCs','SBR paras 5, 10, 11, 13, 29, 32C, 53–57, 101–103'),
 'Asset Liability Management':('All NBFCs incl. HFCs','SBR para 26 + Annex VII; para 89 + Annex XXI (LRM/LCR)'),
 'Income Recognition, Asset Classification and Provisioning':('All NBFCs incl. HFCs','SBR paras 12, 14, 15, 87, 88, 108'),
 'Resolution of Stressed Assets':('All NBFCs (HFCs: specified paras only)','Prudential Framework 2019; IU submissions 2017; Compromise Settlements & Technical Write-offs 2023; Project Finance 2025; Trade Relief 2025'),
 'Treatment of Wilful Defaulters and Large Defaulters':('All NBFCs incl. HFCs','MD Treatment of Wilful Defaulters and Large Defaulters, 2024'),
 'Financial Statements: Presentation and Disclosures':('All NBFCs','SBR paras 30.1, 10, 27; Annex II (excl. para 3); disclosure formats from gold loan, project finance, SSA, TLE directions'),
 'Prudential Norms on Declaration of Dividends':('All NBFCs','SBR para 33'),
 'Know Your Customer':('All NBFCs','Master Direction – KYC Direction, 2016 (as applicable to NBFCs)'),
 'Responsible Business Conduct':('All NBFCs except HFCs','SBR Ch VII, Annex XIII; Digital Lending, KFS, pre-payment charges, gold/silver collateral, microfinance, penal charges instructions'),
 'Managing Risks in Outsourcing':('All NBFCs (BL exempt from Ch IV and IT parts of Ch II); IT outsourcing: new agreements from 28 Nov 2025, existing at renewal or 10 Apr 2026','SBR Annex XIII; MD Outsourcing of IT Services 2023'),
 'Climate Finance and Management of Climate Change Risks':('All NBFCs','Framework for acceptance of green deposits, 2023'),
 'Housing Finance Companies':('All HFCs','MD NBFC–HFC (Reserve Bank) Directions, 2021'),
 'Standalone Primary Dealers':('SPDs','MD Standalone Primary Dealers, 2016'),
 'Core Investment Companies':('All CICs','MD Core Investment Companies (Reserve Bank) Directions, 2016'),
 'Account Aggregator':('NBFC-AA','MD NBFC–Account Aggregator, 2016'),
 'Mortgage Guarantee Companies':('MGCs','MD Mortgage Guarantee Companies, 2016'),
 'Peer to Peer Lending Platform':('NBFC-P2P','MD NBFC–P2P, 2017'),
 'Microfinance Institution':('NBFC-MFIs (microfinance loans by all NBFCs per RBC/Credit Facilities)','MD Regulatory Framework for Microfinance Loans 2022; SBR Annex XXVI'),
 'Non-Operative Financial Holding Companies':('NOFHCs','[VERIFY source instruments]'),
 'Miscellaneous':('NBFC-Factors/ICCs, IDF-NBFCs, BL and ML specifics (per secondary summary)','[VERIFY]')}
SUBJ=[('Governance','Governance & Board'),('Know Your Customer','KYC / AML'),('Responsible Business Conduct','Customer Conduct & Fair Practices'),
 ('Income Recognition','Asset Classification & Provisioning'),('Asset Classification','Asset Classification & Provisioning'),('Capital Adequacy','Capital'),('Capital Charge','Capital'),('Market Risk','Capital'),('Credit Valuation Adjustment','Capital'),
 ('Asset Liability','ALM & Liquidity'),('Cash Reserve Ratio','ALM & Liquidity (CRR/SLR)'),('Outsourcing','Outsourcing & Third-party'),('Credit Information Reporting','Credit Information Reporting'),
 ('Credit Facilities','Credit / Lending'),('Credit Risk Management','Risk Management'),('Concentration Risk','Risk Management'),('Climate','Climate & ESG Risk'),
 ('Transfer and Distribution','Credit Risk Transfer & Co-lending'),('Securitisation','Securitisation'),('Investment Portfolio','Investments'),('Stressed Assets','Stressed Assets & Recovery'),
 ('Wilful Defaulters','Defaulters'),('Financial Statements','Financial Reporting & Disclosure'),('Dividend','Dividends & Profit Distribution'),('Deposits','Deposits'),('Interest Rate','Interest Rates'),
 ('Resource Raising','Funding & Resources'),('Cards','Cards'),('Branch','Branch & Channel Authorisation'),('Digital Banking Channels','Branch & Channel Authorisation'),
 ('Amalgamation','Corporate Actions'),('Acquisition','Ownership & Control'),('Licensing','Licensing & Registration'),('Wholly Owned Subsidiaries','Licensing & Registration'),
 ('Registration','Licensing & Registration'),('Undertaking of Financial Services','Permitted Business'),('Forthcoming','Forthcoming Instructions'),('Miscellaneous Supervisory','Supervisory (Misc incl. PCA)'),
 ('Miscellaneous','Miscellaneous'),('Supervisory Returns','Supervisory Returns'),('Statutory Audit','Statutory Audit'),('Internal Audit','Internal Audit'),('Concurrent Audit','Concurrent Audit'),
 ("Auditor's Report","Auditor's Report"),('Fraud','Fraud Risk Management'),('Digital Payment Security','Cyber & Payment Security'),('Cybersecurity','Cyber & Technology Risk'),('Compliance Function','Compliance Function'),
 ('Internal Ombudsman','Grievance Redressal'),('Housing Finance','Type-specific: HFC'),('Core Investment','Type-specific: CIC'),('Mortgage Guarantee','Type-specific: MGC'),('Primary Dealers','Type-specific: SPD'),
 ('Holding Companies','Type-specific: NOFHC'),('Microfinance','Type-specific: MFI'),('Peer to Peer','Type-specific: P2P'),('Account Aggregator','Type-specific: AA'),
 ('Asset Reconstruction Companies (principal','Type-specific: ARC'),('Credit Information Companies (principal','Type-specific: CIC (credit information)')]
def subject(t):
    for k,v in SUBJ:
        if k.lower() in t.lower(): return v
    return 'Other'
P1=['Governance','Know Your Customer','Responsible Business Conduct','Income Recognition','Capital Adequacy','Asset Liability','Outsourcing','Credit Information Reporting','Credit Facilities',
    'Registration','Compliance Function','Cybersecurity','Fraud','Internal Audit','Supervisory Returns','Statutory Audit',"Auditor's Report",'Internal Ombudsman','Financial Statements',
    'Concentration Risk','Credit Risk Management','Cash Reserve Ratio','Digital Payment Security','Miscellaneous Supervisory','Concurrent Audit']
P3=['Licensing','Forthcoming','Voluntary Amalgamation','Wholly Owned','Digital Banking Channels','Capital Charge for Credit Risk','Market Risk','Credit Valuation']
def priority(t):
    if any(k.lower() in t.lower() for k in P3): return 'P3'
    if any(k.lower() in t.lower() for k in P1): return 'P1'
    return 'P2'
XL=['Know Your Customer','Responsible Business Conduct','Income Recognition','Capital Adequacy','Governance','Credit Facilities','Resolution of Stressed','Housing Finance','Asset Reconstruction Companies (principal','Cybersecurity','Miscellaneous Supervisory','Asset Classification, Provisioning']
L=['Asset Liability','Outsourcing','Investment Portfolio','Financial Statements','Concentration','Credit Risk Management','Transfer and Distribution','Securitisation','Fraud','Internal Audit','Statutory Audit','Supervisory Returns','Compliance Function','Core Investment','Microfinance','Peer to Peer','Account Aggregator','Registration','Acceptance of Public Deposits','Credit Information Companies (principal','Miscellaneous','Digital Payment Security','Concurrent Audit','Market Risk','Capital Charge','Credit Valuation','Climate','Cash Reserve']
S=['Dividend','Amalgamation','Branch Authorisation','Forthcoming','Interest Rates on Advances','Digital Banking Channels','Wholly Owned','Licensing']
def size(t):
    for k in XL:
        if k.lower() in t.lower(): return 'XL'
    for k in S:
        if k.lower() in t.lower(): return 'S'
    for k in L:
        if k.lower() in t.lower(): return 'L'
    return 'M'
SIZE_INFO={'S':('15–40','4–8'),'M':('40–90','8–16'),'L':('90–180','16–32'),'XL':('180–400','32–60')}
OWNER_SHEETS={460,461,462,463,464,465,466,467,468}  # owner's Sep 2026 NBFC DoS sheets
STD_ABBR={'Commercial Banks':'Commercial Banks','Small Finance Banks':'Small Finance Banks','Payments Banks':'Payments Banks','Local Area Banks':'Local Area Banks','Regional Rural Banks':'Regional Rural Banks',
 'Urban Co-operative Banks':'Urban Co-operative Banks','Rural Co-operative Banks':'Rural Co-operative Banks','All India Financial Institutions':'All India Financial Institutions',
 'Non-Banking Financial Companies':'Non-Banking Financial Companies','Asset Reconstruction Companies':'Asset Reconstruction Companies','Credit Information Companies':'Credit Information Companies'}
def full_title(cat,t,kind='DoR',year='2025'):
    if t.endswith('(Guidelines)'):
        base=t.replace(' (Guidelines)','')
        if '–' in base and ('Licensing' in base): return f'Reserve Bank of India ({base}) Guidelines, {year}'
        if cat=='Urban Co-operative Banks': return f'Reserve Bank of India (Urban Co-operative Banks – {base}) Guidelines, {year}'
        return f'Reserve Bank of India ({base}) Guidelines, {year}'
    if '(principal direction)' in t: return f'Reserve Bank of India ({t.replace(" (principal direction)","")}) Directions, {year}'
    if cat=='Non-Banking Financial Companies' and t in ('Housing Finance Companies','Core Investment Companies','Mortgage Guarantee Companies','Standalone Primary Dealers','Non-Operative Financial Holding Companies'):
        return f'Reserve Bank of India ({t}) Directions, {year}'
    if t=='Miscellaneous Supervisory': return f'Reserve Bank of India ({STD_ABBR.get(cat,cat)} – Miscellaneous) Supervisory Directions, {year}'
    return f'Reserve Bank of India ({STD_ABBR.get(cat,cat)} – {re.sub(r" \(2026\)$","",t)}) Directions, {year}'
CONFIRMED_REF={344,345,355,356,363,349,359,273,275,460,461,462,463,464,465,466,467,468,469,415,424,442}
rows=[]
def add(**k): rows.append(k)
for cat,(did,lst) in DOR.items():
    abbr,fam,typ,tpl,cmd,wave=CAT[cat]
    for pid,ser,t,upd in lst:
        ptype=typ; appl='All '+cat+' (read the Direction\'s applicability paragraph; record layer/tier carve-outs in Column C/S)'; sup='Pre-28-Nov-2025 Master Directions/circulars on this subject for '+cat+'; circulars withdrawn by RBI/2025-26/100 (Annex of 9,445 circulars) — build the para concordance in P05'
        if cat=='Non-Banking Financial Companies':
            key=t
            if key in NBFC_TYPE: ptype=NBFC_TYPE[key][0]
            else: ptype='per applicability'
            if key in NBFC_APPL: appl,sup=NBFC_APPL[key]; appl+=' [secondary summary — VERIFY applicability para]' if 'confirmed' not in appl else ''; sup='Consolidates: '+sup+' [secondary summary — VERIFY]'
        ft=full_title(cat,t)
        ref=f'RBI/DOR/2025-26/{ser}'
        basis=('Confirmed on RBI Direction page (25 Oct 2026)' if ser==344 else ('Derived from RBI MD serial; matches number cited in v1.0 research [VERIFY on document]' if ser in CONFIRMED_REF else 'Derived from RBI MD serial (PDF code) [VERIFY on document]'))
        add(cat=cat,abbr=abbr,fam=fam,typ=ptype,tpl=tpl,cmd=cmd,title=ft,short=t,ref=ref,basis=basis,date='28 Nov 2025',upd=upd,dept='DoR',subj=subject(t),appl=appl,sup=sup,
            amend=(f'Updated as on {upd} (RBI page); list amending instruments in P05 [VERIFY]' if upd else 'None shown on RBI page as of 25 Oct 2026'),
            wave=wave,url=SRC+str(did),pid=pid,src='DoR consolidated set 28 Nov 2025',status='Not started',note='')
for cat,pid,ser,t,dt,fy in DOR_LATER:
    abbr,fam,typ,tpl,cmd,wave=CAT[cat]
    yr='2026'; ft=full_title(cat,t,year=yr)
    add(cat=cat,abbr=abbr,fam=fam,typ=typ,tpl=tpl,cmd=cmd,title=ft,short=t,ref=f'RBI/DOR/{fy}/{ser} [VERIFY]',basis='Derived from RBI MD serial [VERIFY on document]',date=dt,upd='',dept='DoR',subj=subject(t),
        appl='All '+cat+(' (supersedes the 2025 dividend direction)' if 'Dividend' in t else ''),sup=('Supersedes 2025 dividend direction of 28 Nov 2025' if 'Dividend' in t and cat!='Regional Rural Banks' else 'New instrument after the consolidation'),
        amend='—',wave=wave,url=SRC+str(DOR[cat][0]),pid=pid,src='DoR after 28 Nov 2025',status='Not started',note='')
for cat,(did,lst) in DOS.items():
    abbr,fam,typ,tpl,cmd,wave=CAT[cat]
    for pid,ser,t in lst:
        ft=full_title(cat,t,year='2026')
        st='Drafted' if (cat=='Non-Banking Financial Companies' and ser in OWNER_SHEETS) else 'Not started'
        note='Owner built this sheet in Sep 2026 — re-check against the "updated as on" text and re-cite' if st=='Drafted' else ''
        add(cat=cat,abbr=abbr,fam=fam,typ=typ,tpl=tpl,cmd=cmd,title=ft,short=t,ref=f'RBI/DoS/2026-27/{ser}',basis=('Derived from RBI MD serial; matches number cited in v1.0 research and RBI press release [VERIFY on document]' if ser in CONFIRMED_REF else 'Derived from RBI MD serial (PDF code) [VERIFY on document]'),
            date='31 Jul 2026',upd='',dept='DoS',subj=subject(t),appl='All '+cat+' unless the Direction carves out layers/tiers (e.g. Compliance Function: NBFC-ML and UL) [VERIFY para]',
            sup='Repeals earlier supervisory instructions on this subject incl. MD on Filing of Supervisory Returns, 2024 (for Supervisory Returns directions) [VERIFY]',
            amend='Check "Updated as on" on RBI page',wave=wave,url=SRC+str(did),pid=pid,src='DoS set 31 Jul 2026',status=st,note=note)
for cat,pid,ser,t in CEPD:
    abbr,fam,typ,tpl,cmd,wave=CAT[cat]
    add(cat=cat,abbr=abbr,fam=fam,typ=typ,tpl=tpl,cmd=cmd,title=full_title(cat if cat!='Non-bank PPI Issuers' else 'Non-Bank Prepaid Payment Instruments Issuers',t,year='2026'),short=t,
        ref=f'RBI/CEPD/2025-26/{ser.rstrip("?")} [VERIFY]',basis='Derived from PDF code; conflicts with v1.0 citation /384 for NBFC — VERIFY on document',date='14 Jan 2026',upd='',dept='CEPD',subj='Grievance Redressal',
        appl='REs within the scope thresholds of the Direction [VERIFY]',sup='Supersedes MD Internal Ombudsman for Regulated Entities, 2023 (for these REs) [VERIFY]',amend='—',wave=wave,url=SRC+'338',pid=pid,src='CEPD 2026',status='Not started',note='')
add(cat='All RBI-regulated REs in scope',abbr='ALL',fam='BANK, UCB, RCB, NBFC, PAY (in scope)',typ='—',tpl='(each family template)',cmd='00_Part2_Master_Command.md + family command',title='Reserve Bank – Integrated Ombudsman Scheme, 2026',short='Integrated Ombudsman Scheme',
    ref='[VERIFY reference no.]',basis='Press release, Jan 2026; effective 1 Jul 2026',date='Jan 2026 [VERIFY date]',upd='',dept='CEPD',subj='Grievance Redressal',appl='Banks, NBFCs, non-bank PPI issuers, CICs as covered [VERIFY coverage of UCBs/RCBs]',
    sup='Supersedes RB-IOS 2021',amend='—',wave='W1',url='https://website.rbi.org.in/web/rbi/-/press-releases/rbi-issues-reserve-bank-integrated-ombudsman-scheme-2026',pid='',src='CEPD scheme',status='Not started',note='Write once per family with family-specific Column C/D values')

# ordering
waveord={'W1':1,'W2':2,'W3':3}
catord={c:i for i,c in enumerate(['Non-Banking Financial Companies','Urban Co-operative Banks','Rural Co-operative Banks','Asset Reconstruction Companies','Credit Information Companies','Commercial Banks','Small Finance Banks','Payments Banks','Local Area Banks','Regional Rural Banks','All India Financial Institutions','Non-bank PPI Issuers','All RBI-regulated REs in scope'])}
for r in rows:
    r['prio']=priority(r['short']); r['size']=size(r['short']); r['rows_est'],r['hours']=SIZE_INFO[r['size']]
    tag=' (DoS 2026)' if r['dept']=='DoS' else (' (CEPD 2026)' if r['dept']=='CEPD' and r['abbr']!='ALL' else '')
    short=r['short'].replace(' (principal direction)','').replace(' (Guidelines)','').replace(':',' -').replace("'", '')
    r['fname']=f"RBI-{r['abbr']} - {short}{tag} - Compliance Questions.xlsx"
rows.sort(key=lambda r:(waveord[r['wave']],catord.get(r['cat'],99),{'P1':1,'P2':2,'P3':3}[r['prio']],r['dept']!='DoR',r['short']))

wb=Workbook()
H=['S.No','RE category','LDC family code (P15)','Entity type (P15)','Template code (Appendix E)','Part-2 command file','Direction title','Reference number','Reference basis','Date of issue','Last updated (RBI page)','Issuing department','Subject area','Applicability notes','Supersedes / consolidated from','Amendments after issue','Priority','Production wave','Size class','Indicative checkpoints','Indicative effort (hours, draft + review)','Proposed question-sheet file name','Question-sheet status','Owner','Reviewer','Uploaded to LDC (date)','RBI page','RBI page id','Source set','Verification status','Notes']
ws=wb.active; ws.title='Catalogue'
hdr_fill=PatternFill('solid',fgColor='0A2540'); hf=Font(bold=True,color='FFFFFF',name='Calibri',size=10)
thin=Side(style='thin',color='B8C2CC'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ws.append(H)
for c in ws[1]: c.fill=hdr_fill; c.font=hf; c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=bd
for i,r in enumerate(rows,1):
    ver='Title, date and RBI page id confirmed on rbi.org.in (25 Oct 2026)' if r['src']!='CEPD scheme' else 'Press release found; reference and coverage to verify'
    ws.append([i,r['cat'],r['fam'],r['typ'],r['tpl'],r['cmd'],r['title'],r['ref'],r['basis'],r['date'],r['upd'],r['dept'],r['subj'],r['appl'],r['sup'],r['amend'],r['prio'],r['wave'],r['size'],r['rows_est'],r['hours'],r['fname'],r['status'],'','','',r['url'],r['pid'],r['src'],ver,r['note']])
widths=[6,22,14,18,18,22,60,24,30,12,14,9,24,40,45,30,8,8,7,11,12,52,14,14,14,14,40,9,22,34,40]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top'); c.font=Font(name='Calibri',size=9); c.border=bd
ws.freeze_panes='H2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(H))}{len(rows)+1}'; ws.row_dimensions[1].height=42
n=len(rows)+1
dv=DataValidation(type='list',formula1='"Not started,In progress,Drafted,Reviewed,Uploaded,Superseded"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'W2:W{n+300}')
dv2=DataValidation(type='list',formula1='"P1,P2,P3"',allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'Q2:Q{n+300}')
dv3=DataValidation(type='list',formula1='"W1,W2,W3"',allow_blank=True); ws.add_data_validation(dv3); dv3.add(f'R2:R{n+300}')

# DoS sheet
ds=wb.create_sheet('DoS_2026'); dh=['S.No','RE category','Direction title','Reference number','Date','Template code','Proposed question-sheet file name','Question-sheet status','RBI page id','Notes']
ds.append(dh)
k=0
for r in rows:
    if r['dept']=='DoS':
        k+=1; ds.append([k,r['cat'],r['title'],r['ref'],r['date'],r['tpl'],r['fname'],r['status'],r['pid'],r['note']])
for i,w in enumerate([6,24,70,24,12,18,60,14,10,50],1): ds.column_dimensions[get_column_letter(i)].width=w
# Superseded
ss=wb.create_sheet('Superseded'); ss.append(['RE category','MD serial (2025-26)','Instrument','Superseded by','Treatment in LDC'])
for c,ser,t,by in SUPERSEDED: ss.append([c,f'RBI/DOR/2025-26/{ser} [VERIFY]',f'Reserve Bank of India ({c} – {t}) Directions, 2025',by,'Keep as historical version in cnt; no new question sheet; checkpoints, if any, re-pointed to the 2026 direction'])
for i,w in enumerate([24,26,70,32,60],1): ss.column_dimensions[get_column_letter(i)].width=w
# summary
sm=wb.create_sheet('By_Category_Summary')
sm.append(['RE category','Template code','Wave','DoR 28 Nov 2025 (listed)','DoR after 28 Nov 2025','DoS 31 Jul 2026','CEPD 2026','Total','P1','P2','P3','Indicative effort (hours, low–high)'])
from collections import defaultdict
agg=defaultdict(lambda: defaultdict(int)); eff=defaultdict(lambda:[0,0]); meta={}
for r in rows:
    a=agg[r['cat']]; meta[r['cat']]=(r['tpl'],r['wave'])
    key={'DoR consolidated set 28 Nov 2025':'d','DoR after 28 Nov 2025':'l','DoS set 31 Jul 2026':'s'}.get(r['src'],'c')
    a[key]+=1; a[r['prio']]+=1; lo,hi=r['hours'].split('–'); eff[r['cat']][0]+=int(lo); eff[r['cat']][1]+=int(hi)
tot=[0]*9
for c in sorted(agg,key=lambda x:(waveord[meta[x][1]],catord.get(x,99))):
    a=agg[c]; vals=[a['d'],a['l'],a['s'],a['c'],a['d']+a['l']+a['s']+a['c'],a['P1'],a['P2'],a['P3']]
    sm.append([c,meta[c][0],meta[c][1]]+vals+[f'{eff[c][0]}–{eff[c][1]}'])
    for i,v in enumerate(vals): tot[i]+=v
sm.append(['TOTAL','','']+tot[:8]+[f"{sum(e[0] for e in eff.values())}–{sum(e[1] for e in eff.values())}"])
for i,w in enumerate([32,20,7,14,14,12,10,8,6,6,6,22],1): sm.column_dimensions[get_column_letter(i)].width=w
# waves
wv=wb.create_sheet('Waves')
for line in [
 ['Wave','Families / RE categories','Why this order','Exit criterion'],
 ['W1','NBFC (35 DoR + 9 DoS + NBFC Internal Ombudsman) + RB-IOS 2026','LDC launch market (master §13 Wave 1); owner already holds NBFC question banks and the 9 DoS sheets of Sep 2026; RBI/DoS/2026-27/460 requires ML/UL NBFCs to use workflow-based compliance tools','All P1 sheets uploaded and published; P2 drafted'],
 ['W2','UCB (27 + 9 DoS), RCB (26 + 3 DoS), ARC (4 + 1 DoS), Credit Information Companies (2 + 2 DoS + IO)','Co-operative and ARC segments are named target markets; return generation for ARC and co-op ships in Wave 2','All P1 sheets uploaded; P2 drafted'],
 ['W3','Commercial Banks (32 + 4 later DoR + 9 DoS + IO), SFB (30 + 1 + 9 + IO), PB (17 + 1 + 9 + IO), LAB (25 + 1 + 7), RRB (23 + 1 + 2), AIFI (19 + 4), non-bank PPI IO','Bank families are longer sales cycles; shared subject structure lets sheets be cloned from the SCB master sheet with category-specific applicability','All sheets uploaded'],
 ['Rule','Within a wave: P1 before P2 before P3; DoR before DoS only where the DoS direction depends on DoR definitions; one subject is written first for the lead category and then cloned and adapted for sibling categories (e.g. KYC: NBFC → UCB → RCB → banks)','Cuts drafting effort for cloned subjects by an estimated 40–60% [REC]','']]:
    wv.append(line)
for i,w in enumerate([8,60,70,40],1): wv.column_dimensions[get_column_letter(i)].width=w
# instructions
ins=wb.create_sheet('Instructions',0)
txt=["LawDocsComply — RBI Directions Catalogue v1.1 (25 October 2026) (v1.1)",
"Purpose: the master list of RBI instruments that will be converted into Part-2 question sheets and uploaded to the LawDocsComply content library. For RBI families the content base is the set of consolidated Master Directions issued by the Department of Regulation on 28 November 2025 — '238 function-wise Master Directions' in RBI's press release, 244 Master Directions in RBI circular RBI/2025-26/100, serials RBI/DOR/2025-26/143 to /386 (240 currently listed; 4 dividend directions superseded in March 2026), their amendments, the new DoR directions issued since, the 64 Department of Supervision directions of 31 July 2026 (RBI/DoS/2026-27/408 to /471), the CEPD Internal Ombudsman directions of 14 January 2026 and the Reserve Bank – Integrated Ombudsman Scheme, 2026.",
"Sources: rbi.org.in Master Directions pages (BS_ViewMasterDirections.aspx?did=403–413 for DoR, 419–429 for DoS, 338 for CEPD), accessed 25 Oct 2026; RBI circular RBI/2025-26/100 (Consolidation of Regulations – Withdrawal of circulars, 28 Nov 2025: 244 Master Directions; 9,445 circulars withdrawn).",
"Reference numbers: RBI's PDF file codes for the 28 Nov 2025 set equal the RBI/DOR/2025-26/NNN reference (checked on NBFC Governance /344 and eight other cited directions). Rows marked 'Derived' must be confirmed on the Direction's first page before the sheet is published (Column I).",
"How to use: filter Column R (wave) and Q (priority); write the sheet with the template in Column E and the command in Column F; save under the name in Column V; update Columns W–Z. When RBI updates a Direction, change Column K and raise a P05 change task.",
"Status values: Not started · In progress · Drafted · Reviewed · Uploaded · Superseded.",
"Size classes (indicative, for planning only): S 15–40 checkpoints · M 40–90 · L 90–180 · XL 180–400. Effort = drafting plus review hours per sheet.",
"Do not treat applicability notes marked [VERIFY] or secondary as final: read the Direction's applicability paragraph and record carve-outs in Columns C and S of the question sheet."]
for t in txt: ins.append([t])
ins.column_dimensions['A'].width=160
for row in ins.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ins['A1'].font=Font(bold=True,size=13,color='0A2540')
for sh in [ds,ss,sm,wv]:
    for c in sh[1]: c.fill=hdr_fill; c.font=hf; c.alignment=Alignment(wrap_text=True)
    sh.freeze_panes='A2'
    for row in sh.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save('/home/claude/ldc_v11/catalogue/LDC_RBI_Directions_Catalogue_v1.1.xlsx')
import json; json.dump(rows,open('/home/claude/ldc_v11/catalogue/_build/rows.json','w'),indent=0,default=str)
print(len(rows)); 
from collections import Counter; print(Counter(r['src'] for r in rows)); print(Counter((r['wave']) for r in rows)); print(Counter(r['prio'] for r in rows))
