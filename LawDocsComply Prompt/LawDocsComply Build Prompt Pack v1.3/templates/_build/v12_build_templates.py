# -*- coding: utf-8 -*-
"""Build LawDocsComply v2.0 templates (pack v1.2, D-49 / D-57).

Outputs (in <out_dir>, default ../ relative to this file):
  LDC-TPL-<FAMILY>-v2.0 - <Name>.xlsx  for the 27 entity families
  LDC-TPL-RISK-v2.0 - Risk Register.xlsx
  LDC-TPL-GOV-v2.0 - Governance Register.xlsx

Layout per the template contract v2.0 (templates/_build/v12_contract.py; Appendix E Part A):
row 1 coloured group bands, row 2 headers, data from row 3; LIST VIEW always visible; other groups
outlined (level 1 = first column of each group, level 2 = the rest), collapsed by default; freeze D3;
AutoFilter on row 2; validations rows 3-5000; Department -> Function -> Sub-function dependent lists;
print area = list view; Lists, Read Me and protected _Template_Info sheets.

Usage: python3 v12_build_templates.py [out_dir]
"""
import hashlib
import os
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.packaging.custom import StringProperty

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import v12_contract as K  # noqa: E402

FONT = 'Calibri'
THIN = Side(style='thin', color='FF9E9E9E')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
HDR_ALIGN = Alignment(wrap_text=True, vertical='center', horizontal='center')
DATE_FMT = 'DD-MMM-YYYY'


def fill(hex6):
    return PatternFill('solid', fgColor='FF' + hex6)


def contract_json(columns, groups):
    import json
    glabel = {g: l for g, l, _ in groups}
    gcol = {g: c for g, _, c in groups}
    role = {'Author': 'AUTHOR', 'Author/User': 'AUTHOR', 'User': 'USER', 'System': 'SYSTEM'}
    return json.dumps([dict(key=c['key'], header=c['header'], group=glabel[c['group']], colour=gcol[c['group']], role=role[c['by']],
                            required_in_content=c['at'] == 'Required', conditional=c['at'] == 'Conditional', list=c['list'], multi_value=c['multi'],
                            list_view=c['group'] == groups[0][0], ext=bool(c.get('ext')), position_hint=i)
                       for i, c in enumerate(columns, 1)], ensure_ascii=False, separators=(',', ':'))


def header_hash(headers):
    return hashlib.sha256(K.SEP.join(headers).encode('utf-8')).hexdigest()


def key_hash(keys):
    return hashlib.sha256(K.SEP.join(keys).encode('utf-8')).hexdigest()


class _WholeColumns:
    """Print area expressed as whole columns (e.g. 'Sheet'!$A:$I), written verbatim to _xlnm.Print_Area."""

    def __init__(self, text):
        self.text = text

    def __str__(self):
        return self.text

    def __bool__(self):
        return True


# ------------------------------------------------------------------ Lists sheet
def write_lists(wb, prefix, lists, fmatrix=None):
    ls = wb.create_sheet('Lists')
    names = {}
    c = 1
    bold = Font(name=FONT, bold=True, size=10)
    for key, vals in lists.items():
        L = get_column_letter(c)
        ls.cell(row=1, column=c, value=key).font = bold
        ls.cell(row=1, column=c).fill = fill('D9E1F2')
        for i, v in enumerate(vals, 2):
            ls.cell(row=i, column=c, value=v)
        nm = f'{prefix}_{key}'
        wb.defined_names[nm] = DefinedName(nm, attr_text=f'Lists!${L}$2:${L}${len(vals) + 1}')
        names[key] = nm
        ls.column_dimensions[L].width = min(48, max(12, max(len(str(v)) for v in vals) + 2))
        c += 1
    if fmatrix:
        fmap, smap = fmatrix
        for label, mp, top, hdr in (('FUNCTIONS by Department →', fmap, 'FuncTop', 'FuncHdr'),
                                    ('SUB-FUNCTIONS by Function →', smap, 'SubTop', 'SubHdr')):
            c += 1  # spacer column
            ls.cell(row=1, column=c, value=label).font = Font(name=FONT, bold=True, italic=True, size=10)
            c += 1
            first = c
            for head, vals in mp.items():
                ls.cell(row=1, column=c, value=head).font = bold
                ls.cell(row=1, column=c).fill = fill('FCE4D6')
                for i, v in enumerate(vals, 2):
                    ls.cell(row=i, column=c, value=v)
                ls.column_dimensions[get_column_letter(c)].width = 26
                c += 1
            last = c - 1
            fl, ll = get_column_letter(first), get_column_letter(last)
            wb.defined_names[f'{prefix}_{top}'] = DefinedName(f'{prefix}_{top}', attr_text=f'Lists!${fl}$2')
            wb.defined_names[f'{prefix}_{hdr}'] = DefinedName(f'{prefix}_{hdr}', attr_text=f'Lists!${fl}$1:${ll}$1')
            names[top], names[hdr] = f'{prefix}_{top}', f'{prefix}_{hdr}'
    ls.freeze_panes = 'A2'
    ls.sheet_properties.tabColor = 'FF8EA9DB'
    return names


# ------------------------------------------------------------------ main register sheet
def write_register(wb, sheet_name, columns, groups, names, prefix, code, dept_letter='E'):
    ws = wb.active
    ws.title = sheet_name
    colour = {g: c for g, _, c in groups}
    label = {g: l for g, l, _ in groups}
    letters = {}
    for i, c in enumerate(columns, 1):
        letters[c['key']] = get_column_letter(i)
    last = get_column_letter(len(columns))

    # row 1 bands (merged per group) and row 2 headers
    start = 1
    for i in range(1, len(columns) + 2):
        g_prev = columns[start - 1]['group']
        if i == len(columns) + 1 or columns[i - 1]['group'] != g_prev:
            cell = ws.cell(row=1, column=start, value=label[g_prev])
            for j in range(start, i):
                ws.cell(row=1, column=j).fill = fill(colour[g_prev])
                ws.cell(row=1, column=j).border = BOX
            cell.font = Font(name=FONT, bold=True, size=11)
            cell.alignment = Alignment(horizontal='left', vertical='center')
            if i - 1 > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i - 1)
            start = i
    for i, c in enumerate(columns, 1):
        cell = ws.cell(row=2, column=i, value=c['header'])
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = fill(colour[c['group']])
        cell.alignment = HDR_ALIGN
        cell.border = BOX
        tip = f"[{c['key']}] {c['by']} · {c['at']}\n{c['rule']}"
        cm = Comment(tip, 'LawDocsComply')
        cm.width, cm.height = 320, 140
        cell.comment = cm
        ws.column_dimensions[get_column_letter(i)].width = c['width']
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 48

    # data-row formats
    for i, c in enumerate(columns, 1):
        for r in range(K.FIRST_DATA_ROW, K.FORMAT_LAST_ROW + 1):
            cell = ws.cell(row=r, column=i)
            cell.alignment = WRAP_TOP
            if c['list'] == 'date':
                cell.number_format = DATE_FMT

    # outline groups: LIST VIEW level 0; first column of each later group level 1; remaining level 2; all hidden
    seen = set()
    for i, c in enumerate(columns, 1):
        L = get_column_letter(i)
        cd = ws.column_dimensions[L]
        if c['group'] == groups[0][0]:
            continue
        lvl = 1 if c['group'] not in seen else 2
        seen.add(c['group'])
        cd.outline_level = lvl
        cd.hidden = True
    # collapsed marker on the summary column to the right of each collapsed run
    for i, c in enumerate(columns, 1):
        nxt = columns[i]['group'] if i < len(columns) else None
        if c['group'] != groups[0][0] and nxt != c['group']:
            ws.column_dimensions[get_column_letter(i + 1)].collapsed = True
    ws.column_dimensions[get_column_letter(len(columns) + 1)].width = 3
    ws.sheet_format.outlineLevelCol = 2
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.sheet_properties.outlinePr.summaryBelow = True

    ws.freeze_panes = 'D3'
    ws.auto_filter.ref = f'A2:{last}2'

    # data validation
    dvs = []
    for i, c in enumerate(columns, 1):
        L = get_column_letter(i)
        rng = f'{L}{K.FIRST_DATA_ROW}:{L}{K.DV_LAST_ROW}'
        lst = c['list']
        if not lst:
            continue
        if lst == 'date':
            dv = DataValidation(type='date', operator='between', formula1='DATE(1950,1,1)', formula2='DATE(2100,12,31)', allow_blank=True,
                                showErrorMessage=True, errorTitle='Date', error='Enter a date (DD-MMM-YYYY) between 1950 and 2100.',
                                showInputMessage=True, promptTitle=c['header'][:32], prompt='Date, DD-MMM-YYYY')
        elif lst == 'FUNC':
            f = (f'OFFSET({names["FuncTop"]},0,MATCH(${dept_letter}3,{names["FuncHdr"]},0)-1,'
                 f'COUNTA(OFFSET({names["FuncTop"]},0,MATCH(${dept_letter}3,{names["FuncHdr"]},0)-1,60,1)),1)')
            assert len(f) < 255, len(f)
            dv = DataValidation(type='list', formula1=f, allow_blank=True, showErrorMessage=True, errorStyle='warning',
                                errorTitle='Department Function', error='Choose a function listed for the Department in column E.',
                                showInputMessage=True, promptTitle='Department Function', prompt='Choose the Department (column E) first.')
        elif lst == 'SUBF':
            fl = letters['dept_function']
            f = (f'OFFSET({names["SubTop"]},0,MATCH(${fl}3,{names["SubHdr"]},0)-1,'
                 f'COUNTA(OFFSET({names["SubTop"]},0,MATCH(${fl}3,{names["SubHdr"]},0)-1,60,1)),1)')
            assert len(f) < 255, len(f)
            dv = DataValidation(type='list', formula1=f, allow_blank=True, showErrorMessage=True, errorStyle='warning',
                                errorTitle='Department Sub-function', error='Choose a sub-function listed for the Function.',
                                showInputMessage=True, promptTitle='Department Sub-function', prompt=f'Choose the Department Function (column {fl}) first.')
        else:
            multi = c['multi']
            soft = c.get('soft', False)
            dv = DataValidation(type='list', formula1=names[lst], allow_blank=True, showErrorMessage=not multi,
                                errorStyle='warning' if soft else 'stop', errorTitle=c['header'][:32],
                                error=('Not in the list. Click Yes only if no list value fits (the importer will flag it).' if soft else 'Choose a value from the list.'),
                                showInputMessage=True, promptTitle=c['header'][:32],
                                prompt=('Choose from the list; several values may be typed separated by "; "' if multi else 'Choose from the list')[:255])
        dv.add(rng)
        ws.add_data_validation(dv)
        dvs.append((c['key'], lst))

    # conditional formatting on RAG (and Days to Due) if present
    if 'rag' in letters:
        rl = letters['rag']
        targets = [f'{rl}{K.FIRST_DATA_ROW}:{rl}{K.DV_LAST_ROW}']
        for k2 in ('days_to_due', 'days_to_review'):
            if k2 in letters:
                targets.append(f'{letters[k2]}{K.FIRST_DATA_ROW}:{letters[k2]}{K.DV_LAST_ROW}')
        for t in targets:
            for nm, hx in K.RAG_FILL.items():
                ws.conditional_formatting.add(t, FormulaRule(formula=[f'${rl}{K.FIRST_DATA_ROW}="{nm}"'], fill=fill(hx), stopIfTrue=True))

    # print setup: list view only
    ws._print_area = _WholeColumns(f"'{sheet_name}'!$A:$I")  # openpyxl only accepts cell ranges; Excel accepts whole columns
    ws.print_title_rows = '1:2'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = code
    ws.oddFooter.center.text = '&F'
    ws.oddFooter.right.text = 'Page &P of &N'
    ws.oddFooter.left.size = ws.oddFooter.center.size = ws.oddFooter.right.size = 8
    ws.sheet_properties.tabColor = 'FFFFC000'
    ws.sheet_view.zoomScale = 90
    return ws, letters, dvs


# ------------------------------------------------------------------ Read Me
def write_readme(wb, code, title, columns, groups, sheet_name, extra_lines, ext_cols):
    rm = wb.create_sheet('Read Me')
    rm.column_dimensions['A'].width = 30
    rm.column_dimensions['B'].width = 110
    big = Font(name=FONT, bold=True, size=14)
    bold = Font(name=FONT, bold=True, size=11)
    norm = Font(name=FONT, size=10)
    r = 1
    rm.cell(row=r, column=1, value=f'{title} — how to use this template ({code})').font = big
    r += 2

    def para(a, b='', fa=bold, fl=None):
        nonlocal r
        ca = rm.cell(row=r, column=1, value=a)
        ca.font = fa
        ca.alignment = WRAP_TOP
        cb = rm.cell(row=r, column=2, value=b)
        cb.font = norm
        cb.alignment = WRAP_TOP
        if fl:
            ca.fill = fill(fl)
        r += 1

    para('What this is', f'The LawDocsComply v2.0 template ({code}). One workbook = one instrument (master direction / regulation / Act / standard). '
         f'Write one actionable per row in the "{sheet_name}" sheet, starting at row 3. Do not rename sheets, edit rows 1–2, or move columns.')
    para('Compact view', 'The workbook opens in LIST VIEW: only the yellow columns are visible, so the list fits one screen and prints on A4 landscape. '
         'Use the outline buttons above the column letters: "1" = list view only; "2" = list view + the first column of every group; '
         '"3" = all columns. The "+" / "−" buttons above each group open or close that group alone.')
    para('Filter and sort', 'Row 2 has AutoFilter on every column: use the arrow in any header to filter or sort. Rows 1–2 and columns A–C stay frozen.')
    para('Header tips', 'Hover over any header (row 2) to see its machine key, who fills it and the rule. The full contract is in _Template_Info.')
    r += 1
    para('Column groups and colours', '', bold)
    for g, lab, colr in groups:
        cols = [c['header'] for c in columns if c['group'] == g]
        para(lab, f'{len(cols)} columns: ' + ' · '.join(cols), Font(name=FONT, bold=True, size=10), colr)
    r += 1
    if ext_cols:
        para('Family extension columns', 'Inside APPLICABILITY this family adds: ' + ' · '.join(c['header'] for c in ext_cols) +
             '. Leave blank where the actionable does not depend on that dimension.')
    for a, b in extra_lines:
        para(a, b)
    r += 1
    para('RAG colours (exports)', 'Red F8696B = overdue / critical window · Amber FFC000 = due soon · Green 63BE7B = on track · Grey BFBFBF = not applicable / exempt / withdrawn · '
         'Blue 9BC2E6 = event-based item awaiting its trigger. RAG is computed by LawDocsComply on export (published rules); do not type it.')
    para('Who fills what', 'Author columns are filled by the content writer (Part-2). User columns are filled by makers, checkers and auditors in the tool '
         '(or in a working copy) and are ignored when content is imported. System columns are written by LawDocsComply on export.')
    para('Dates', 'Enter real dates (DD-MMM-YYYY). Text such as "31.03.2026" is rejected by the importer.')
    para('Several values', 'Where a column allows several values, separate them with "; " (semicolon + space).')
    rm.sheet_properties.tabColor = 'FF70AD47'
    return rm


# ------------------------------------------------------------------ _Template_Info
def write_info(wb, code, fam_code, fam_name, regulator, columns, groups, names, dvs, letters, extra_ident, instruments=None, example=None):
    info = wb.create_sheet('_Template_Info')
    bold = Font(name=FONT, bold=True, size=10)
    norm = Font(name=FONT, size=10)
    band = fill('D9E1F2')
    widths = [6, 8, 18, 34, 24, 12, 12, 18, 7, 80]
    for i, w in enumerate(widths, 1):
        info.column_dimensions[get_column_letter(i)].width = w
    headers = [c['header'] for c in columns]
    keys = [c['key'] for c in columns]
    r = 1
    info.cell(row=r, column=1, value=f'LawDocsComply template identity and column contract — {fam_name}').font = Font(name=FONT, bold=True, size=13)
    r = 3
    ident = [('template_code', code), ('template_version', K.CONTRACT_VERSION), ('contract_version', K.CONTRACT_VERSION),
             ('family_code', fam_code), ('family_name', fam_name), ('regulator', regulator), ('built_on', K.BUILD_DATE),
             ('data_sheet', columns and wb.active.title), ('header_row', K.HEADER_ROW), ('first_data_row', K.FIRST_DATA_ROW),
             ('validation_last_row', K.DV_LAST_ROW), ('column_count', len(columns)),
             ('extension_keys', '; '.join(c['key'] for c in columns if c.get('ext'))),
             ('header_hash', header_hash(headers)), ('key_hash', key_hash(keys)),
             ('hash_method', 'SHA-256 over the UTF-8 bytes of the row-2 header texts (NFC, trimmed) — or of the machine keys — in column order, joined with U+001F'),
             ('generated_at', K.BUILD_DATE + 'T00:00:00+05:30'),
             ('contract_json', contract_json(columns, groups)),
             ('multi_value_delimiter', K.MULTI_SEP), ('named_ranges', ', '.join(sorted(names.values()))),
             ('importer_contract', 'Appendix E Part A (contract v2.0), LawDocsComply Build Prompt Pack v1.2; error codes LDC-TPL-nnn'),
             ] + extra_ident
    info.cell(row=r - 1, column=1, value='IDENTITY').font = bold
    for k, v in ident:
        info.cell(row=r, column=1, value=k).font = bold
        info.cell(row=r, column=1).fill = band
        info.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c2 = info.cell(row=r, column=4, value=v)
        c2.font = norm
        c2.alignment = WRAP_TOP
        info.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
        r += 1
    r += 1
    info.cell(row=r, column=1, value='COLUMN CONTRACT').font = bold
    r += 1
    heads = ['#', 'Column', 'Group', 'Header (row 2, exact)', 'Key', 'Filled by', 'At authoring', 'List / validation', 'Multi', 'Rule']
    for i, h in enumerate(heads, 1):
        c = info.cell(row=r, column=i, value=h)
        c.font = bold
        c.fill = band
        c.border = BOX
    r += 1
    glabel = {g: l for g, l, _ in groups}
    for i, c in enumerate(columns, 1):
        lst = c['list']
        lst_txt = '' if not lst else ('date' if lst == 'date' else ('dependent: Function by Department' if lst == 'FUNC' else
                                                                    ('dependent: Sub-function by Function' if lst == 'SUBF' else names.get(lst, lst))))
        row = [i, get_column_letter(i), glabel[c['group']], c['header'], c['key'], c['by'], c['at'], lst_txt, 'Yes' if c['multi'] else '', c['rule']]
        for j, v in enumerate(row, 1):
            cell = info.cell(row=r, column=j, value=v)
            cell.font = norm
            cell.alignment = WRAP_TOP
            cell.border = BOX
        r += 1
    info.cell(row=r, column=1, value='END CONTRACT').font = bold
    r += 2
    if instruments:
        info.cell(row=r, column=1, value='INSTRUMENTS THIS FAMILY COVERS (examples; verify current versions before writing)').font = bold
        r += 1
        for it in instruments:
            info.cell(row=r, column=1, value='• ' + it).font = norm
            info.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
            r += 1
        r += 1
    if example:
        info.cell(row=r, column=1, value='WORKED EXAMPLE (illustrative; carried from v1.0 via the legacy map — verify the citation before use)').font = bold
        r += 1
        for k2, v in example:
            info.cell(row=r, column=1, value=k2).font = bold
            info.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            cc = info.cell(row=r, column=4, value=v)
            cc.font = norm
            cc.alignment = WRAP_TOP
            info.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
            r += 1
        r += 1
    info.cell(row=r, column=1, value='LEGACY v1.0 → v2.0 MAP (importer, Appendix E Part A §A.9)').font = bold
    r += 1
    for j, h in enumerate(['v1.0 col', 'v1.0 header', '', 'v2.0 key', 'Note'], 1):
        if h:
            info.cell(row=r, column=j if j < 4 else (4 if j == 4 else 5), value=h).font = bold
    r += 1
    for a, b, k2, note in K.LEGACY_MAP:
        info.cell(row=r, column=1, value=a).font = norm
        info.cell(row=r, column=2, value=b).font = norm
        info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        info.cell(row=r, column=4, value=k2).font = norm
        info.cell(row=r, column=5, value=note).font = norm
        info.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
        r += 1
    info.protection.sheet = True
    info.protection.password = K.INFO_PASSWORD
    info.sheet_properties.tabColor = 'FF7F7F7F'
    return info


def finish(wb, code, out_path, title):
    wb.properties.title = title
    wb.properties.creator = 'LawDocsComply'
    wb.properties.subject = code
    wb.properties.keywords = f'{code}; contract {K.CONTRACT_VERSION}'
    wb.custom_doc_props.append(StringProperty(name='LDC_TemplateCode', value=code))
    wb.custom_doc_props.append(StringProperty(name='LDC_ContractVersion', value=K.CONTRACT_VERSION))
    # sheet order: register, Lists, Read Me, _Template_Info
    order = [wb.active.title, 'Lists', 'Read Me', '_Template_Info']
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(out_path)


# ------------------------------------------------------------------ family templates
def example_pairs(fam):
    """Map the v1.0 worked example (letters) to v2.0 headers through the legacy map."""
    ex = fam.get('example') or {}
    hdr = {c['key']: c['header'] for c in K.BASE_COLUMNS}
    out = []
    m = {'B': 'sno', 'C': 'applic_layer', 'D': 'applic_type', 'F': 'short_title', 'G': 'question', 'H': 'frequency', 'K': 'clause_ref',
         'L': 'explanation', 'M': 'department', 'N': 'related_docs', 'O': 'compliance_type', 'P': 'filing_details', 'Q': 'consequences',
         'R': 'risk_level', 'S': 'conditions', 'U': 'evidence', 'V': 'control_activity'}
    for k in K.BASE_KEYS:
        for L, key in m.items():
            if key == k and L in ex:
                out.append((hdr[k], ex[L]))
    return out


def build_family(fam, out_dir):
    code = f"LDC-TPL-{fam['code']}-v{K.CONTRACT_VERSION}"
    prefix = fam['prefix']
    columns = K.family_columns(fam['code'])
    lists = K.family_lists(fam)
    fmatrix = K.function_matrix(fam['M'])
    wb = openpyxl.Workbook()
    wb.active.title = K.SHEET
    names = write_lists(wb, prefix, lists, fmatrix)
    ws, letters, dvs = write_register(wb, K.SHEET, columns, K.GROUPS, names, prefix, code)
    assert letters['department'] == 'E'
    ext_cols = [c for c in columns if c.get('ext')]
    extra = [('Department → Function → Sub-function', f'Choose Department (column E) first; Department Function (column {letters["dept_function"]}) and '
              f'Department Sub-function (column {letters["dept_subfunction"]}) then offer only the values that belong to it.'),
             ('CompID', 'Leave CompID blank. LawDocsComply assigns a unique, permanent CompID on import (format <AUTH>-<FAM>-<INST>-<NNNN>, e.g., '
              f'{K.COMPID_EXAMPLE[fam["code"]]}). '
              'If you re-import an exported sheet, keep the CompIDs: they tell the importer which items to update.'),
             ('Version and update dates', 'Fill "Name of Master Direction / Regulation", "Instrument Reference No." and "Issue Date of Instrument" on every row. '
              '"Instrument Last Amended On" repeats the issue date if the instrument was never amended. "Actionable Last Updated On" is the date of the regulator change that last affected THAT row, so it differs row to row; '
              'set "Change Status" (New / Amended / Unchanged / Re-cited / Retired) and, for Amended / Re-cited / Retired rows, the previous wording in "Change Note".'),
             ('Event-based items', 'Where Due Date & Frequency is Event Based or Immediately, choose the Event Trigger and state the time limit in Due Date Rule.')]
    write_readme(wb, code, fam['name'], columns, K.GROUPS, K.SHEET, extra, ext_cols)
    write_info(wb, code, fam['code'], fam['name'], fam['regulator'], columns, K.GROUPS, names, dvs, letters,
               [('compid_pattern', ('<issuing AUTH>' if fam['code'] in K.CROSS_AUTHORITY else K.AUTH_CODE[K.family_group(fam["code"])]) + f'-{K.FAMILY_SHORT_CODE[fam["code"]]}-<INST>-<NNNN>; tenant T<TENANT4>-<INST>-<NNNN>'),
                ('list_view_columns', 'A:I'), ('legacy_accepted', 'LDC-TPL-*-v1.0 (30 columns A–AD) via the legacy map'),
                ('part2_command', f"part2_prompts/00_Part2_Master_Command.md + the {fam['code']} family command")],
               fam.get('instruments'), example_pairs(fam))
    name = f"LDC-TPL-{fam['code']}-v{K.CONTRACT_VERSION} - {K.SHORT[fam['code']]}.xlsx"
    finish(wb, code, os.path.join(out_dir, name), f"LawDocsComply template — {fam['name']}")
    return name


def dept_union():
    seen = []
    for f in K.all_families():
        for d in f['M']:
            if d not in seen:
                seen.append(d)
    return seen


def build_risk(out_dir):
    code = f'LDC-TPL-RISK-v{K.CONTRACT_VERSION}'
    wb = openpyxl.Workbook()
    names = write_lists(wb, 'RISK', K.risk_lists(dept_union()))
    ws, letters, dvs = write_register(wb, 'Risk Register', K.RISK_COLUMNS, K.RISK_GROUPS, names, 'RISK', code)
    extra = [('RiskID', 'Leave RiskID blank. LawDocsComply assigns RSK-<TENANT4>-<NNNNN> (tenant risk) or RSKL-<TAXO>-<NNNN> (library risk); '
              'the ID never changes and is never reused. KRIs and controls carry KRI- and CTL- IDs.'),
             ('Ratings', 'Inherent and residual ratings follow the tenant risk matrix (likelihood × impact on 1–5 scales). '
              'Link the CompIDs whose non-compliance would crystallise the risk so the risk heat map and the compliance RAG stay connected.')]
    write_readme(wb, code, 'Risk Register', K.RISK_COLUMNS, K.RISK_GROUPS, 'Risk Register', extra, [])
    write_info(wb, code, 'RISK', 'Risk Register', 'Tenant risk function (library: LawDocsComply)', K.RISK_COLUMNS, K.RISK_GROUPS, names, dvs, letters,
               [('id_pattern', 'RSK-<TENANT4>-<NNNNN> | RSKL-<TAXO>-<NNNN>'), ('list_view_columns', 'A:I')])
    name = f'LDC-TPL-RISK-v{K.CONTRACT_VERSION} - Risk Register.xlsx'
    finish(wb, code, os.path.join(out_dir, name), 'LawDocsComply template — Risk Register')
    return name


def build_gov(out_dir):
    code = f'LDC-TPL-GOV-v{K.CONTRACT_VERSION}'
    wb = openpyxl.Workbook()
    names = write_lists(wb, 'GOV', K.gov_lists())
    ws, letters, dvs = write_register(wb, 'Governance Register', K.GOV_COLUMNS, K.GOV_GROUPS, names, 'GOV', code)
    extra = [('GovID', 'Leave GovID blank. LawDocsComply assigns GOVL-<AUTH>-<NNNN> (library requirement) or GOV-<TENANT4>-<TYPE>-<NNNNN> '
              '(tenant item: MTG meeting, AGD agenda item, MIN minutes, RES resolution, REG register entry, FIL filing, POL policy, COM committee, '
              'DIR director event, KMP KMP event, DEC declaration). Fill "GovID Type" so the right TYPE segment is used.')]
    write_readme(wb, code, 'Governance Register', K.GOV_COLUMNS, K.GOV_GROUPS, 'Governance Register', extra, [])
    write_info(wb, code, 'GOV', 'Governance Register', 'Companies Act / Secretarial Standards / LODR / sectoral governance rules', K.GOV_COLUMNS,
               K.GOV_GROUPS, names, dvs, letters, [('id_pattern', 'GOVL-<AUTH>-<NNNN> | GOV-<TENANT4>-<TYPE>-<NNNNN>'), ('list_view_columns', 'A:I')])
    name = f'LDC-TPL-GOV-v{K.CONTRACT_VERSION} - Governance Register.xlsx'
    finish(wb, code, os.path.join(out_dir, name), 'LawDocsComply template — Governance Register')
    return name


def main():
    out_dir = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(HERE)
    os.makedirs(out_dir, exist_ok=True)
    built = [build_family(f, out_dir) for f in K.all_families()]
    built += [build_risk(out_dir), build_gov(out_dir)]
    for b in built:
        print('built', b)
    print(len(built), 'workbooks')


if __name__ == '__main__':
    main()
