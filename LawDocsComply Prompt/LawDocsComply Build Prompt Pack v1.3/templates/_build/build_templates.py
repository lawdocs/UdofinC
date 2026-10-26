# -*- coding: utf-8 -*-
"""Build LawDocsComply entity-family question templates (LDC-TPL-<FAMILY>-v1.0).

Base: the owner's IFSCA question sheet (Appointed Actuary Regulations 2022), which carries the
most complete IFSCA vocabulary. Rows 1-2 of Sheet1 are kept byte-identical; data rows are cleared;
Rough is rebuilt with family lists; named ranges are family-prefixed; a protected _Template_Info
sheet carries the template identity, column contract, hashes and Part-2 instructions.

Usage: python3 build_templates.py <base_xlsx> <out_dir>
"""
import copy
import hashlib
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.packaging.custom import StringProperty

sys.path.insert(0, __file__.rsplit('/', 1)[0])
from family_vocab import FAMILIES  # noqa: E402

VERSION = '1.0'
BUILD_DATE = '2026-10-25'
DV_LAST_ROW = 2000
FORMAT_LAST_ROW = 1000
SEP = '\x1f'  # unit separator used for header hashing

COLS = [get_column_letter(i) for i in range(1, 31)]  # A..AD

# Column contract: col -> (owner, at_upload, rule)
CONTRACT = {
    'A': ('System', 'Blank', 'Checkbox used by the software. Leave blank.'),
    'B': ('Content writer', 'Required', 'Serial number 1..n, integer, unique and contiguous within the sheet.'),
    'C': ('Content writer', 'Required', 'Applicability layer / category. One value from the Column C list (multiple values separated by "; " only where the obligation applies to several).'),
    'D': ('Content writer', 'Required', 'Applicable entity type(s). One or more values from the Column D list separated by "; " (semicolon + space).'),
    'E': ('System', 'Blank', 'CompID. Assigned by the software on import (stable obligation ID). Leave blank; on re-export it is filled and must not be edited.'),
    'F': ('Content writer', 'Required', 'Short title: the heading of the para / regulation, or a short descriptive title (max 200 characters).'),
    'G': ('Content writer', 'Required', 'Compliance question / actionable, phrased as a closed question ("Has the <entity> ...?"). One actionable per row; split bundled sub-clauses into separate rows.'),
    'H': ('Content writer', 'Required', 'Due date & frequency. One value from the Column H list. Put the exact due-date rule in Column P after "Due:".'),
    'I': ('Maker (user)', 'Blank', 'Date of compliance. Filled by the user in the software.'),
    'J': ('System / Maker', 'Blank', 'Status. Set by the workflow in the software.'),
    'K': ('Content writer', 'Required', 'Exact clause / para / regulation / section reference, including "as substituted / inserted vide <notification> dated <date>" where the text was amended.'),
    'L': ('Content writer', 'Required', 'Interpretation / explanation in plain language: quote the provision, explain each element, note links to other provisions.'),
    'M': ('Content writer', 'Required', 'Department. One value from the Column M list (tenants may rename departments in the software).'),
    'N': ('Content writer', 'Optional', 'Related document / compliance / weblinks: instrument title and date | official URL | related instruments | INCORPORATES: amendments. Parts separated by " | ".'),
    'O': ('Content writer', 'Required', 'Compliance type. One value from the Column O list.'),
    'P': ('Content writer', 'Required', 'Filing / reporting / intimation. Start with a value from the Column P list ("Not Required" if none); optionally add " — " and detail, and "Due: <rule>" for the due-date rule.'),
    'Q': ('Content writer', 'Required', 'Consequences of non-compliance / penalty / impact: the provision breached, the enforcement power (Act and section) and the practical consequence.'),
    'R': ('Content writer', 'Required', 'Risk categorisation. High / Medium / Low per the calibration rule in the Part-2 command.'),
    'S': ('Content writer', 'Required', 'Conditional applicability / exemptions / thresholds / carve-outs. Write "Applies to all <entities>." if none.'),
    'T': ('System', 'Blank', 'Maker user / e-mail ID. Filled by the software from assignments.'),
    'U': ('Content writer', 'Required', 'Evidence to upload: the specific artefact(s) that prove compliance (document, register, acknowledgement, URL with archive).'),
    'V': ('Content writer', 'Required', 'Control activity: imperative, step-by-step instruction to the maker on what to do and what to produce.'),
    'W': ('Maker (user)', 'Blank', 'Status update by maker. User-filled in the software.'),
    'X': ('Maker (user)', 'Blank', 'Comments by maker. User-filled.'),
    'Y': ('System', 'Blank', 'Checker user / e-mail ID. Filled by the software.'),
    'Z': ('Checker (user)', 'Blank', 'Filing / compliance date. User-filled.'),
    'AA': ('Checker (user)', 'Blank', 'Review: Approve / Reject / Reallocate. Set by the workflow.'),
    'AB': ('Checker (user)', 'Blank', 'Comments by checker. User-filled.'),
    'AC': ('System', 'Blank', 'Assigned auditor / e-mail ID. Filled by the software.'),
    'AD': ('Auditor (user)', 'Blank', 'Comments by auditor. User-filled.'),
}

SHORT = {'RBI-SCB': 'Commercial Banks SFBs PBs LABs RRBs', 'RBI-UCB': 'Urban Co-operative Banks', 'RBI-RCB': 'Rural Co-operative Banks',
         'RBI-NBFC': 'NBFCs', 'RBI-ARC': 'Asset Reconstruction Companies', 'RBI-OTH': 'Other RBI Regulated Entities',
         'COOP-CCS': 'Credit Co-operative Societies', 'COOP-GEN': 'Other Co-operative Societies', 'IFSCA-GEN': 'IFSCA Regulated Entities',
         'IFSCA-FCFU': 'IFSCA Finance Companies and Units', 'IFSCA-FME': 'IFSCA Fund Management Entities', 'SEBI-MF': 'Mutual Funds',
         'SEBI-AIF': 'Alternative Investment Funds', 'SEBI-LISTED': 'Listed Entities', 'SEBI-REIT-INVIT': 'REITs and InvITs',
         'SEBI-INT': 'SEBI Intermediaries', 'IRDAI-INS': 'Insurers', 'IRDAI-INT': 'Insurance Intermediaries', 'MCA-CO': 'Companies',
         'MCA-LLP': 'LLPs', 'DATA-PRIV': 'Data Protection and Privacy', 'CYBER': 'Cyber Security', 'LAB-POSH': 'Labour Codes and POSH',
         'TAX-GST': 'GST and Direct Tax'}

VAL_COLS = {'H': 'Freq', 'C': 'Sector', 'D': 'Entity', 'M': 'Dept', 'O': 'CompType', 'R': 'Risk', 'P': 'Filing'}


def header_hash(ws, row):
    vals = ['' if ws[f'{c}{row}'].value is None else str(ws[f'{c}{row}'].value) for c in COLS]
    return hashlib.sha256(SEP.join(vals).encode('utf-8')).hexdigest()


def build(base_path, out_dir, fam):
    wb = openpyxl.load_workbook(base_path)
    ws = wb['Sheet1']
    rough = wb['Rough']

    # capture data-row style from row 3 of the base
    row_style = {c: (copy.copy(ws[f'{c}3'].font), copy.copy(ws[f'{c}3'].border), copy.copy(ws[f'{c}3'].alignment),
                     copy.copy(ws[f'{c}3'].fill), ws[f'{c}3'].number_format) for c in COLS}
    hdr_font = copy.copy(rough['A1'].font)
    item_font = copy.copy(rough['A2'].font)

    # clear data rows
    if ws.max_row >= 3:
        ws.delete_rows(3, ws.max_row - 2)
    for r in range(3, FORMAT_LAST_ROW + 1):
        for c in COLS:
            cell = ws[f'{c}{r}']
            f, b, a, fi, nf = row_style[c]
            cell.font, cell.border, cell.alignment, cell.fill, cell.number_format = f, b, a, fi, nf
    for r in range(3, FORMAT_LAST_ROW + 1):
        ws.row_dimensions[r].height = None

    # rebuild Rough (IFSCA-GEN keeps the owner's Rough byte-identical)
    prefix = fam['prefix']
    sections = [('Due Date & Frequency (Column H)', fam['H'], 'Freq'), (fam['c_label'], fam['C'], 'Sector'), (fam['d_label'], fam['D'], 'Entity'),
                ('Department (Column M)', fam['M'], 'Dept'), ('Compliance Type (Column O)', fam['O'], 'CompType'), ('Risk (Column R)', fam['R'] if 'R' in fam else ['High', 'Medium', 'Low'], 'Risk'),
                ('Filing/Reporting (Column P - reference)', fam['P'], 'Filing')]
    ranges = {}
    if fam.get('exact_owner'):
        # locate existing sections in owner's Rough
        vals = [c.value for c in rough['A']]
        i = 0
        order = ['Freq', 'Sector', 'Entity', 'Dept', 'CompType', 'Risk', 'Filing']
        k = 0
        while i < len(vals) and k < len(order):
            if vals[i] is not None and (i == 0 or vals[i - 1] is None):
                start = i + 2  # 1-based row of first item
                j = i + 1
                while j < len(vals) and vals[j] is not None:
                    j += 1
                ranges[order[k]] = (start, j)  # j is 1-based last row (index j-1)
                k += 1
                i = j
            i += 1
        for key, exp in [('Freq', fam['H']), ('Sector', fam['C']), ('Entity', fam['D']), ('Dept', fam['M']), ('CompType', fam['O']), ('Filing', fam['P'])]:
            got = [rough[f'A{r}'].value for r in range(ranges[key][0], ranges[key][1] + 1)]
            assert got == exp, (key, got, exp)
    else:
        rough.delete_rows(1, rough.max_row)
        r = 1
        for label, items, key in sections:
            rough[f'A{r}'] = label
            rough[f'A{r}'].font = copy.copy(hdr_font)
            start = r + 1
            for it in items:
                r += 1
                rough[f'A{r}'] = it
                rough[f'A{r}'].font = copy.copy(item_font)
            ranges[key] = (start, r)
            r += 2
        rough.column_dimensions['A'].width = 70

    # defined names
    for n in list(wb.defined_names.keys()):
        del wb.defined_names[n]
    for key, (a, b) in ranges.items():
        nm = f'{prefix}_{key}'
        wb.defined_names[nm] = DefinedName(nm, attr_text=f'Rough!$A${a}:$A${b}')

    # data validation
    ws.data_validations.dataValidation = []
    for col, key in VAL_COLS.items():
        dv = DataValidation(type='list', formula1=f'{prefix}_{key}', allow_blank=True, showErrorMessage=False, showInputMessage=True,
                            promptTitle=f'Column {col}', prompt=('Choose from the list; several values may be typed separated by "; "' if col in ('C', 'D') else
                                                                 ('Start with a list value; you may add " — detail" and "Due: <rule>"' if col == 'P' else 'Choose from the list')))
        dv.add(f'{col}3:{col}{DV_LAST_ROW}')
        ws.add_data_validation(dv)

    # _Template_Info
    if '_Template_Info' in wb.sheetnames:
        del wb['_Template_Info']
    info = wb.create_sheet('_Template_Info')
    code = f"LDC-TPL-{fam['code']}-v{VERSION}"
    h2 = header_hash(ws, 2)
    h1 = header_hash(ws, 1)
    bold = Font(name='Times New Roman', bold=True, size=11)
    norm = Font(name='Times New Roman', size=10)
    title = Font(name='Times New Roman', bold=True, size=14)
    band = PatternFill('solid', fgColor='FFD9E1F2')
    thin = Side(style='thin', color='FF999999')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical='top')
    info.column_dimensions['A'].width = 30
    info.column_dimensions['B'].width = 26
    info.column_dimensions['C'].width = 34
    info.column_dimensions['D'].width = 18
    info.column_dimensions['E'].width = 14
    info.column_dimensions['F'].width = 90

    def put(r, c, v, font=norm, fill=None, border=True):
        cell = info.cell(row=r, column=c, value=v)
        cell.font = font
        cell.alignment = wrap
        if fill:
            cell.fill = fill
        if border:
            cell.border = box
        return cell

    r = 1
    put(r, 1, f'LawDocsComply Question Template — {fam["name"]}', title, border=False)
    r += 2
    ident = [('Template code', code), ('Template version', VERSION), ('Family code', fam['code']), ('Family name', fam['name']),
             ('Regulator / authority', fam['regulator']), ('Built on', BUILD_DATE), ('Base layout', "Owner's IFSCA question sheet (Appointed Actuary Regulations 2022 sample; canonical vocabulary per master X-38); rows 1–2 and columns A–AD unchanged"),
             ('Data sheet', 'Sheet1 (do not rename)'), ('List sheet', 'Rough (named ranges ' + ', '.join(f'{prefix}_{k}' for k in ranges) + ')'),
             ('Row-2 header SHA-256', h2), ('Row-1 header SHA-256', h1), ('Hash method', 'SHA-256 over UTF-8 of the 30 cell values A..AD joined with U+001F; empty cell = empty string'),
             ('Multi-value delimiter (C, D)', '"; " (semicolon + space)'), ('Data rows', 'Start at row 3. Validations cover rows 3–2000; split larger instruments into Part A/Part B files.'),
             ('Importer contract', 'Appendix E, LawDocsComply Build Prompt Pack v1.0 (error codes LDC-TPL-nnn)'),
             ('Part-2 command', f"part2_prompts/00_Part2_Master_Command.md + the {fam['code']} family command")]
    for k, v in ident:
        put(r, 1, k, bold, band)
        info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        put(r, 2, v)
        r += 1

    r += 1
    put(r, 1, 'Instruments this family covers (examples; verify current versions before writing)', bold, border=False)
    r += 1
    for it in fam['instruments']:
        info.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        put(r, 1, '• ' + it)
        r += 1

    r += 1
    put(r, 1, 'Column contract', bold, border=False)
    r += 1
    for c, h in enumerate(['Column', 'Row-1 group', 'Row-2 header', 'Filled by', 'At upload', 'Rule'], 1):
        put(r, c, h, bold, band)
    r += 1
    group = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1:
            for ci in range(rng.min_col, rng.max_col + 1):
                group[get_column_letter(ci)] = ws.cell(row=1, column=rng.min_col).value
    group.setdefault('X', 'Maker')
    for col in COLS:
        owner, at, rule = CONTRACT[col]
        for ci, v in enumerate([col, group.get(col, ''), ws[f'{col}2'].value, owner, at, rule], 1):
            put(r, ci, v)
        r += 1

    r += 1
    put(r, 1, 'Instructions for Part 2 (writing questions)', bold, border=False)
    r += 1
    steps = ['1. Use this file unchanged as the starting point: do not edit rows 1–2, do not rename sheets, do not add or move columns.',
             '2. One workbook per instrument (master direction / regulation / circular / guideline / Act). File name: "<Family code> - <Instrument short title> <Year>.xlsx" (add " vN" for revisions; an existing numbered sector prefix such as "01 Banking - " may be kept).',
             '3. Write one actionable per row; split paragraphs that bundle sub-clauses (policy contents, committee functions, report matters) into separate rows.',
             '4. Fill B, C, D, F, G, H, K, L, M, N, O, P, Q, R, S, U and V. Leave A, E, I, J, T and W–AD blank.',
             '5. Use the dropdown values exactly; where several entity types apply, type them in Column D separated by "; ".',
             '6. Put the due-date rule in Column P after "Due:" (e.g., "Due: within 15 days of the end of each quarter").',
             '7. Cite the exact para / regulation in Column K, including the amending notification where text was substituted or inserted.',
             '8. Follow part2_prompts/00_Part2_Master_Command.md, the family command and 99_QA_Checklist before upload.',
             '9. Upload in LawDocsComply: Content Studio > Import Question Sheet. The importer checks the template code, header hash, vocabularies and required columns and shows a preview before anything is published.']
    for s in steps:
        info.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        put(r, 1, s)
        r += 1

    r += 1
    put(r, 1, 'Worked example row (illustrative only — do not copy into Sheet1 without verifying the citation)', bold, border=False)
    r += 1
    for c, h in enumerate(['Column', 'Row-2 header', 'Example value'], 1):
        put(r, c, h, bold, band)
    info.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    r += 1
    exr = fam['example']
    for col in COLS:
        if col in exr:
            put(r, 1, col)
            put(r, 2, ws[f'{col}2'].value)
            info.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            put(r, 3, exr[col])
            r += 1

    info.protection.sheet = True
    info.protection.password = 'LDC-TPL'
    info.sheet_properties.tabColor = 'FF1F4E79'

    # workbook properties
    wb.properties.creator = 'LawDocsComply'
    wb.properties.title = code
    wb.properties.subject = fam['name']
    wb.properties.keywords = f"LawDocsComply; question template; {fam['code']}"
    wb.properties.lastModifiedBy = 'LawDocsComply'
    for k, v in [('LDC_TemplateCode', code), ('LDC_TemplateVersion', VERSION), ('LDC_Family', fam['code']), ('LDC_HeaderSHA256', h2)]:
        try:
            wb.custom_doc_props.append(StringProperty(name=k, value=v))
        except Exception:
            pass
    wb.active = 0
    readable = SHORT.get(fam['code'], fam['name']).replace('/', '-').replace(':', '-')
    out = f"{out_dir}/{code} - {readable}.xlsx"
    wb.save(out)
    return out, code, h2, {k: len(v) for k, v in [('C', fam['C']), ('D', fam['D']), ('H', fam['H']), ('M', fam['M']), ('O', fam['O']), ('P', fam['P'])]}


if __name__ == '__main__':
    base, out_dir = sys.argv[1], sys.argv[2]
    for fam in FAMILIES:
        fam.setdefault('R', ['High', 'Medium', 'Low'])
        out, code, h2, counts = build(base, out_dir, fam)
        print(code, counts, out.rsplit('/', 1)[1])
