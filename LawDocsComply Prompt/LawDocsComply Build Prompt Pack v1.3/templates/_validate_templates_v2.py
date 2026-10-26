# -*- coding: utf-8 -*-
"""Validate LawDocsComply v2.0 templates (contract v2.0, pack v1.2).

Usage: python3 _validate_templates_v2.py [templates_dir] [report_path]
Checks every LDC-TPL-*-v2.0*.xlsx in templates_dir against templates/_build/v12_contract.py and writes a
Markdown report (default: templates_dir/_validation_report_v2.md). Exit code 0 only if all workbooks pass.
The v1.0 validator (_validate_templates.py) is kept for the legacy files in _legacy_v1.0/.
"""
import glob
import hashlib
import os
import re
import sys
import warnings
import zipfile

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

warnings.filterwarnings('ignore', message='Print area cannot be set')
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '_build'))
import v12_contract as K  # noqa: E402
from family_vocab import DEPT_BASE, COMP_BASE, FREQ_BASE  # noqa: E402

FAM = {f['code']: f for f in K.all_families()}


def colmap(ws):
    """Expand openpyxl column dimension ranges -> {index: (hidden, outline_level)}."""
    out = {}
    for key, cd in ws.column_dimensions.items():
        lo = cd.min or column_index_from_string(key)
        hi = cd.max or lo
        for i in range(lo, hi + 1):
            out[i] = (bool(cd.hidden), int(cd.outline_level or 0))
    return out


def defined_names_xml(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read('xl/workbook.xml').decode('utf-8')
    return dict(re.findall(r'<definedName name="([^"]+)"[^>]*>([^<]+)</definedName>', xml))


def read_info(ws):
    ident, contract, r = {}, [], 1
    rows = list(ws.iter_rows(values_only=True))
    i = 0
    while i < len(rows):
        row = rows[i]
        if row[0] == 'COLUMN CONTRACT':
            i += 2
            while i < len(rows) and rows[i][0] != 'END CONTRACT':
                rr = rows[i]
                contract.append(dict(n=rr[0], letter=rr[1], group=rr[2], header=rr[3], key=rr[4], by=rr[5], at=rr[6], lst=rr[7], multi=rr[8]))
                i += 1
            break
        if row[0] and row[3] is not None and i >= 2:
            ident[row[0]] = row[3]
        i += 1
    return ident, contract


def validate(path):
    fname = os.path.basename(path)
    m = re.match(r'LDC-TPL-(.+)-v2\.0 - (.+)\.xlsx$', fname)
    res = []

    def chk(tid, ok, msg=''):
        res.append((tid, bool(ok), msg))

    if not m:
        return fname, [('T00', False, 'file name does not follow LDC-TPL-<FAMILY>-v2.0 - <Name>.xlsx')]
    fam_code = m.group(1)
    kind = 'family' if fam_code in FAM else fam_code
    if kind == 'family':
        columns, groups, sheet = K.family_columns(fam_code), K.GROUPS, K.SHEET
    elif kind == 'RISK':
        columns, groups, sheet = K.RISK_COLUMNS, K.RISK_GROUPS, 'Risk Register'
    elif kind == 'GOV':
        columns, groups, sheet = K.GOV_COLUMNS, K.GOV_GROUPS, 'Governance Register'
    else:
        return fname, [('T00', False, f'unknown family {fam_code}')]
    code = f'LDC-TPL-{fam_code}-v2.0'
    wb = openpyxl.load_workbook(path)
    # T01 sheets
    chk('T01', wb.sheetnames == [sheet, 'Lists', 'Read Me', '_Template_Info'], f'sheets {wb.sheetnames}')
    ws, ls, info = wb[sheet], wb['Lists'], wb['_Template_Info']
    ident, contract = read_info(info)
    # T02 identity
    chk('T02', ident.get('template_code') == code and str(ident.get('template_version')) == '2.0' and ident.get('family_code') == fam_code
        and ident.get('data_sheet') == sheet and int(ident.get('column_count', 0)) == len(columns),
        f"identity {ident.get('template_code')} / {ident.get('template_version')} / {ident.get('family_code')} / {ident.get('column_count')}")
    # T03 headers
    hdr = [ws.cell(row=2, column=i).value for i in range(1, len(columns) + 2)]
    exp = [c['header'] for c in columns]
    chk('T03', hdr[:-1] == exp and hdr[-1] is None, 'row-2 headers differ from contract' if hdr[:-1] != exp else '')
    h = hashlib.sha256(K.SEP.join(exp).encode('utf-8')).hexdigest()
    keys_exp = [c['key'] for c in columns]
    kh = hashlib.sha256(K.SEP.join(c['key'] for c in columns).encode('utf-8')).hexdigest()
    import json
    try:
        cj = json.loads(ident.get('contract_json') or '[]')
    except ValueError:
        cj = []
    chk('T04', ident.get('header_hash') == h and ident.get('key_hash') == kh and [c['key'] for c in cj] == keys_exp
        and sum(1 for c in cj if c['list_view']) == 9, 'header/key hash or contract_json mismatch')
    chk('T05', [(c['header'], c['key']) for c in contract] == [(c['header'], c['key']) for c in columns]
        and [c['letter'] for c in contract] == [get_column_letter(i) for i in range(1, len(columns) + 1)], 'contract table differs')
    # T06 bands
    colour = {g: c for g, _, c in groups}
    label = {g: l for g, l, _ in groups}
    merged = {(r.min_col, r.max_col) for r in ws.merged_cells.ranges if r.min_row == 1}
    ok, msg, start = True, '', 1
    order = []
    for i in range(1, len(columns) + 2):
        if i == len(columns) + 1 or columns[i - 1]['group'] != columns[start - 1]['group']:
            g = columns[start - 1]['group']
            order.append(g)
            if ws.cell(row=1, column=start).value != label[g]:
                ok, msg = False, f'band label at col {start}'
            if i - 1 > start and (start, i - 1) not in merged:
                ok, msg = False, f'band {g} not merged'
            for j in range(start, i):
                if (ws.cell(row=1, column=start).fill.fgColor.rgb or '')[-6:] != colour[g] or (ws.cell(row=2, column=j).fill.fgColor.rgb or '')[-6:] != colour[g]:
                    ok, msg = False, f'fill colour col {j}'
            start = i
    chk('T06', ok and order == [g for g, _, _ in groups], msg or f'group order {order}')
    # T07 base keys and extensions
    keys = [c['key'] for c in columns]
    if kind == 'family':
        base_ok = [k for k in keys if not k.startswith('ext_')] == K.BASE_KEYS
        ext = [c for c in columns if c.get('ext')]
        allowed = [e[0] for e in K.EXT[fam_code]]
        pos = keys.index('applic_to')
        ext_ok = len(ext) <= 4 and [c['key'] for c in ext] == allowed and keys[pos + 1:pos + 1 + len(ext)] == allowed and all(c['group'] == 'APPLIC' for c in ext)
        chk('T07', base_ok and ext_ok and ident.get('extension_keys', '') in ('; '.join(allowed), None if not allowed else '; '.join(allowed)),
            f'base_ok={base_ok} ext_ok={ext_ok}')
    else:
        chk('T07', len(set(keys)) == len(keys), 'duplicate keys')
    # T08 lists and vocabulary
    lists = {}
    for ci in range(1, ls.max_column + 1):
        head = ls.cell(row=1, column=ci).value
        if not head:
            continue
        vals = []
        r = 2
        while ls.cell(row=r, column=ci).value is not None:
            vals.append(ls.cell(row=r, column=ci).value)
            r += 1
        lists.setdefault(head, vals)
    if kind == 'family':
        f = FAM[fam_code]
        exp_lists = K.family_lists(f)
        vocab_ok = all(lists.get(k) == v for k, v in exp_lists.items())
        base_ok = (set(FREQ_BASE) <= set(lists['Freq']) and set(DEPT_BASE) <= set(lists['Dept']) and set(COMP_BASE) <= set(lists['CompType'])
                   and lists['Risk'] == ['High', 'Medium', 'Low'] and lists['Status'] == K.STATUS and lists['RiskType'] == K.RISK_TYPE
                   and lists['Change'] == K.CHANGE and lists['RAG'] == K.RAG and lists['AuditResult'] == K.AUDIT_RESULT)
        chk('T08', vocab_ok and base_ok, f'vocab_ok={vocab_ok} base_ok={base_ok}')
    else:
        exp_lists = K.risk_lists([]) if kind == 'RISK' else K.gov_lists()
        chk('T08', all(lists.get(k) == v for k, v in exp_lists.items() if k != 'Dept') and (kind != 'RISK' or len(lists.get('Dept', [])) >= 9), 'list mismatch')
    # T09 named ranges resolve
    prefix = FAM[fam_code]['prefix'] if kind == 'family' else kind
    dn = wb.defined_names
    nr_ok, nr_msg = True, ''
    for k, v in exp_lists.items():
        nm = f'{prefix}_{k}'
        if nm not in dn:
            nr_ok, nr_msg = False, f'missing {nm}'
            continue
        ref = dn[nm].attr_text
        mm = re.match(r"Lists!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)$", ref)
        if not mm:
            nr_ok, nr_msg = False, f'bad ref {nm}={ref}'
            continue
        vals = [ls[f'{mm.group(1)}{r}'].value for r in range(int(mm.group(2)), int(mm.group(4)) + 1)]
        if vals != (v if k != 'Dept' or kind != 'RISK' else vals):
            nr_ok, nr_msg = False, f'{nm} values differ'
    chk('T09', nr_ok, nr_msg)
    # T10 data validation coverage
    dvs = {}
    for d in ws.data_validations.dataValidation:
        for rng in str(d.sqref).split():
            dvs[rng] = d
    dv_ok, dv_msg = True, ''
    for i, c in enumerate(columns, 1):
        L = get_column_letter(i)
        d = dvs.get(f'{L}3:{L}{K.DV_LAST_ROW}')
        lst = c['list']
        if not lst:
            if d is not None:
                dv_ok, dv_msg = False, f'unexpected DV on {c["key"]}'
            continue
        if d is None:
            dv_ok, dv_msg = False, f'no DV rows 3-{K.DV_LAST_ROW} on {c["key"]} ({L})'
            continue
        if lst == 'date':
            if d.type != 'date':
                dv_ok, dv_msg = False, f'{c["key"]} not date DV'
        elif lst == 'FUNC':
            if not (d.type == 'list' and 'OFFSET' in d.formula1 and '$E3' in d.formula1 and f'{prefix}_FuncHdr' in d.formula1):
                dv_ok, dv_msg = False, 'dependent Function DV wrong'
        elif lst == 'SUBF':
            fl = get_column_letter(keys.index('dept_function') + 1)
            if not (d.type == 'list' and f'${fl}3' in d.formula1 and f'{prefix}_SubHdr' in d.formula1):
                dv_ok, dv_msg = False, 'dependent Sub-function DV wrong'
        else:
            if not (d.type == 'list' and d.formula1 == f'{prefix}_{lst}'):
                dv_ok, dv_msg = False, f'{c["key"]} DV formula {d.formula1}'
            if c['multi'] and d.showErrorMessage:
                dv_ok, dv_msg = False, f'{c["key"]} multi-value column must not block'
            if c.get('soft') and not c['multi'] and d.errorStyle != 'warning':
                dv_ok, dv_msg = False, f'{c["key"]} soft list must use a warning alert'
            if not c.get('soft') and not c['multi'] and not (d.showErrorMessage and d.errorStyle in (None, 'stop')):
                dv_ok, dv_msg = False, f'{c["key"]} closed list must use a stop alert'
    chk('T10', dv_ok, dv_msg)
    # T11 dependent matrix completeness (families)
    if kind == 'family':
        fh = [c.value for c in ls[1]]
        fm, sm = K.function_matrix(FAM[fam_code]['M'])
        ref_f = dn[f'{prefix}_FuncHdr'].attr_text
        a, b = re.match(r'Lists!\$([A-Z]+)\$1:\$([A-Z]+)\$1', ref_f).groups()
        heads_f = [ls.cell(row=1, column=i).value for i in range(column_index_from_string(a), column_index_from_string(b) + 1)]
        ref_s = dn[f'{prefix}_SubHdr'].attr_text
        a2, b2 = re.match(r'Lists!\$([A-Z]+)\$1:\$([A-Z]+)\$1', ref_s).groups()
        heads_s = [ls.cell(row=1, column=i).value for i in range(column_index_from_string(a2), column_index_from_string(b2) + 1)]
        all_fn = {fn for v in fm.values() for fn in v}
        chk('T11', heads_f == lists['Dept'] and all_fn <= set(heads_s) and len(heads_s) == len(set(heads_s)), 'function matrix incomplete')
    else:
        chk('T11', True, 'n/a')
    # T12 outline / visibility
    cm = colmap(ws)
    out_ok, out_msg, seen = True, '', set()
    for i, c in enumerate(columns, 1):
        hid, lvl = cm.get(i, (False, 0))
        if c['group'] == groups[0][0]:
            if hid or lvl:
                out_ok, out_msg = False, f'list column {i} hidden/outlined'
        else:
            want = 1 if c['group'] not in seen else 2
            seen.add(c['group'])
            if not hid or lvl != want:
                out_ok, out_msg = False, f'col {get_column_letter(i)} hidden={hid} level={lvl} want {want}'
    chk('T12', out_ok and len([c for c in columns if c['group'] == groups[0][0]]) == 9, out_msg)
    # T13 freeze, autofilter
    chk('T13', ws.freeze_panes == 'D3' and ws.auto_filter.ref == f'A2:{get_column_letter(len(columns))}2', f'{ws.freeze_panes} {ws.auto_filter.ref}')
    # T14 print set-up
    dnx = defined_names_xml(path)
    pa, pt = dnx.get('_xlnm.Print_Area', ''), dnx.get('_xlnm.Print_Titles', '')
    chk('T14', pa.endswith('!$A:$I') and pt.endswith('!$1:$2') and ws.page_setup.orientation == 'landscape' and str(ws.page_setup.fitToWidth) == '1'
        and str(ws.page_setup.paperSize) == '9', f'print {pa} {pt}')
    # T15 protection, no data rows, RAG CF, Read Me
    no_data = all(v is None for row in ws.iter_rows(min_row=3, max_row=min(ws.max_row, K.FORMAT_LAST_ROW), values_only=True) for v in row)
    rag_l = get_column_letter(keys.index('rag') + 1)
    cf = [str(r.sqref) for r in ws.conditional_formatting]
    readme_ok = any(wb['Read Me'].cell(row=r, column=1).value == groups[-1][1] for r in range(1, 60))
    chk('T15', info.protection.sheet and no_data and any(s.startswith(f'{rag_l}3') for s in cf) and readme_ok,
        f'protected={info.protection.sheet} no_data={no_data} rag_cf={cf[:1]} readme={readme_ok}')
    return fname, res


COMPID_RE = re.compile(r'^(?P<auth>[A-Z]{2,6})-(?P<fam>[A-Z0-9]{2,8})-(?P<inst>[A-Z0-9]{2,8})-(?P<seq>[0-9]{4,6})$')


def validate_export(path, fam_code='RBI-NBFC'):
    """Check a Working Register export (Appendix E §A.11) against the family contract."""
    res = []
    wb = openpyxl.load_workbook(path)
    columns = K.family_columns(fam_code)
    ws = wb[K.SHEET]
    keys = [c['key'] for c in columns]
    res.append(('X01', [ws.cell(row=2, column=i).value for i in range(1, len(columns) + 1)] == [c['header'] for c in columns], 'headers'))
    lists = K.family_lists(FAM[fam_code])
    fm, sm = K.function_matrix(FAM[fam_code]['M'])
    bad, ids = [], []
    for r in range(3, ws.max_row + 1):
        row = {k: ws.cell(row=r, column=i).value for i, k in enumerate(keys, 1)}
        if row['sno'] is None:
            continue
        ids.append(row['comp_id'])
        for c in columns:
            v, lst = row[c['key']], c['list']
            if v is None or lst in (None, 'date', 'FUNC', 'SUBF') or c.get('soft'):
                continue
            vals = str(v).split(K.MULTI_SEP) if c['multi'] else [v]
            if any(x not in lists[lst] for x in vals) and not (c['key'] == 'filing_details'):
                bad.append(f"row {r} {c['key']}={v}")
        if row['dept_function'] not in fm[row['department']] or (row['dept_subfunction'] and row['dept_subfunction'] not in sm[row['dept_function']]):
            bad.append(f'row {r} function/sub-function')
        if row['frequency'] in ('Event Based', 'Immediately') and not row['event_trigger']:
            bad.append(f'row {r} event trigger missing')
        if not (row['instrument_issue_date'] <= row['question_updated_date'] <= row['instrument_amended_date']):
            bad.append(f'row {r} date logic')
    res.append(('X02', not bad, '; '.join(bad[:5])))
    res.append(('X03', len(ids) == len(set(ids)) and all(COMPID_RE.match(i or '') for i in ids), 'CompID grammar / uniqueness'))
    need = {'Summary', 'Compliance List (grouped)', K.SHEET, 'Lists', 'Read Me', '_Template_Info', '_Export_Info'}
    res.append(('X04', need <= set(wb.sheetnames) and len(wb.sheetnames) > len(need), f'sheets {wb.sheetnames}'))
    rag_l = get_column_letter(keys.index('rag') + 1)
    res.append(('X05', any(str(r.sqref).startswith(f'{rag_l}3') for r in ws.conditional_formatting), 'RAG conditional formatting'))
    return res, len(ids)


def main():
    d = sys.argv[1] if len(sys.argv) > 1 else HERE
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(d, '_validation_report_v2.md')
    files = sorted(glob.glob(os.path.join(d, 'LDC-TPL-*-v2.0 - *.xlsx')))
    results = [validate(p) for p in files]
    npass = sum(all(ok for _, ok, _ in r) for _, r in results)
    tests = {'T01': 'Sheets and order (register, Lists, Read Me, _Template_Info)', 'T02': 'Template identity in _Template_Info (code, version 2.0, family, sheet, column count)',
             'T03': 'Row-2 headers equal the contract, in order, nothing after the last column', 'T04': 'Header SHA-256 and key SHA-256 recorded and correct',
             'T05': 'Column contract table in _Template_Info matches headers, keys and letters', 'T06': 'Row-1 group bands: labels, merges, colours, order',
             'T07': 'Base keys complete and in order; extension columns = family allowance (0–4) placed after Applicable Until',
             'T08': 'Lists: family vocabularies exact; base vocabularies (frequency, department, compliance type, risk, status, risk type, change, RAG, audit result) present',
             'T09': 'Named ranges exist and resolve to the Lists values', 'T10': 'Data validation on every dropdown/date column for rows 3–5000; dependent lists; multi-value columns non-blocking',
             'T11': 'Department → Function → Sub-function matrix complete', 'T12': 'Outline: 9 list-view columns visible; other groups hidden (level 1 first column, level 2 rest)',
             'T13': 'Freeze panes D3; AutoFilter on row 2 across all columns', 'T14': 'Print area A:I, titles rows 1–2, A4 landscape, fit to one page wide',
             'T15': '_Template_Info protected; no data rows; RAG conditional formatting; Read Me legend'}
    lines = ['# Template validation report — contract v2.0 (pack v1.2)', '',
             f'Run: `python3 templates/_validate_templates_v2.py templates` · Contract: `templates/_build/v12_contract.py` · Date: {K.BUILD_DATE}', '',
             f'**Result: {npass}/{len(results)} PASS**', '', '## Tests', '', '| Test | Check |', '|---|---|']
    lines += [f'| {k} | {v} |' for k, v in tests.items()]
    lines += ['', '## Results by workbook', '', '| Workbook | Columns | ' + ' | '.join(tests) + ' | Result |', '|---|---|' + '---|' * len(tests) + '---|']
    for fname, res in results:
        wb = openpyxl.load_workbook(os.path.join(d, fname), read_only=True)
        ncols = sum(1 for v in next(wb.worksheets[0].iter_rows(min_row=2, max_row=2, values_only=True)) if v)
        cells = ' | '.join('✓' if ok else '✗' for _, ok, _ in res)
        lines.append(f'| {fname} | {ncols} | {cells} | {"PASS" if all(ok for _, ok, _ in res) else "FAIL"} |')
    fails = [(f, t, m) for f, r in results for t, ok, m in r if not ok]
    if fails:
        lines += ['', '## Failures', ''] + [f'- {f} — {t}: {m}' for f, t, m in fails]
    sample = os.path.join(d, 'samples', 'SAMPLE - Working Register Export - RBI-NBFC.xlsx')
    if os.path.exists(sample):
        xres, n = validate_export(sample)
        lines += ['', '## Sample export check (synthetic)', '', f'`samples/SAMPLE - Working Register Export - RBI-NBFC.xlsx` — {n} rows', '',
                  '| Check | Meaning | Result |', '|---|---|---|']
        xm = {'X01': 'Register headers equal the RBI-NBFC v2.0 contract', 'X02': 'Every closed-list value valid; Function/Sub-function dependent on Department; Event Trigger present for Event Based / Immediately; issue ≤ actionable updated ≤ last amended',
              'X03': 'CompIDs match the platform grammar and are unique', 'X04': 'Summary, grouped list, per-department sheets, Lists, Read Me, _Template_Info and _Export_Info present',
              'X05': 'RAG conditional formatting present'}
        lines += [f'| {t} | {xm[t]} | {"PASS" if ok else "FAIL: " + m} |' for t, ok, m in xres]
        if not all(ok for _, ok, _ in xres):
            print('SAMPLE FAIL', [x for x in xres if not x[1]])
    lines += ['', '## Legacy', '', 'The 27 v1.0 workbooks are kept in `templates/_legacy_v1.0/` for the legacy importer path (Appendix E §A.9). '
              'They are validated by the unchanged v1.0 validator (`templates/_validate_templates.py`); see `_legacy_v1.0/_validation_report_v1.0.md`.']
    open(out, 'w').write('\n'.join(lines) + '\n')
    print(f'{npass}/{len(results)} PASS')
    for f, t, m in fails:
        print('FAIL', f, t, m)
    sys.exit(0 if npass == len(results) else 1)


if __name__ == '__main__':
    main()
