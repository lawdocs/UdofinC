# -*- coding: utf-8 -*-
"""Build the SYNTHETIC v1.2 export samples for RBI-NBFC (D-50, D-54; Appendix E §A.11):

  templates/samples/SAMPLE - Working Register Export - RBI-NBFC.xlsx
  templates/samples/SAMPLE - Compliance List Export - RBI-NBFC.pdf

All rows, people, tenant and dates are SYNTHETIC test data (export date 25-Oct-2026). Instrument titles and
reference numbers are taken from the RBI Directions Catalogue v1.1 so the Source & Version group looks real,
but the questions, clause numbers and outcomes are illustrative only and are NOT legal content.
Usage: python3 v12_build_samples.py
"""
import math
import re
import os
import sys
import warnings
from collections import Counter, OrderedDict
from datetime import date, timedelta

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore', message='Print area cannot be set')
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import v12_contract as K  # noqa: E402
from v12_build_templates import _WholeColumns  # noqa: E402

TODAY = date(2026, 10, 25)
TEMPLATE = os.path.join(os.path.dirname(HERE), 'LDC-TPL-RBI-NBFC-v2.0 - NBFCs.xlsx')
OUT = os.path.join(os.path.dirname(HERE), 'samples')
TENANT = 'Synthetic Finance Ltd (SYNTHETIC)'
EXPORT_ID = 'EXP-SYNF-2026-10-25-0007'

INST = {  # code: (title, ref, issue, last amended)  — from catalogue v1.1
    'GOV25': ('Reserve Bank of India (Non-Banking Financial Companies – Governance) Directions, 2025', 'RBI/DOR/2025-26/344', date(2025, 11, 28), date(2026, 6, 24)),
    'KYC25': ('Reserve Bank of India (Non-Banking Financial Companies – Know Your Customer) Directions, 2025', 'RBI/DOR/2025-26/361', date(2025, 11, 28), date(2025, 12, 29)),
    'SR26': ('Reserve Bank of India (Non-Banking Financial Companies – Supervisory Returns) Directions, 2026', 'RBI/DoS/2026-27/466', date(2026, 7, 31), None),
    'CRM25': ('Reserve Bank of India (Non-Banking Financial Companies – Credit Risk Management) Directions, 2025', 'RBI/DOR/2025-26/350', date(2025, 11, 28), date(2026, 7, 1)),
    'FRM26': ('Reserve Bank of India (Non-Banking Financial Companies – Fraud Risk Management) Directions, 2026', 'RBI/DoS/2026-27/463', date(2026, 7, 31), None),
    'OUTS25': ('Reserve Bank of India (Non-Banking Financial Companies – Managing Risks in Outsourcing) Directions, 2025', 'RBI/DOR/2025-26/363', date(2025, 11, 28), None),
    'RBC25': ('Reserve Bank of India (Non-Banking Financial Companies – Responsible Business Conduct) Directions, 2025', 'RBI/DOR/2025-26/362', date(2025, 11, 28), date(2026, 7, 1)),
    'CYB26': ('Reserve Bank of India (Non-Banking Financial Companies – Cybersecurity, Technology: Risk, Resilience and Assurance Framework) Directions, 2026', 'RBI/DoS/2026-27/461', date(2026, 7, 31), None),
    'CMPF26': ('Reserve Bank of India (Non-Banking Financial Companies – Compliance Function) Directions, 2026', 'RBI/DoS/2026-27/460', date(2026, 7, 31), None),
    'IRACP25': ('Reserve Bank of India (Non-Banking Financial Companies – Income Recognition, Asset Classification and Provisioning) Directions, 2025', 'RBI/DOR/2025-26/356', date(2025, 11, 28), date(2026, 10, 1)),
}
URL = 'https://www.rbi.org.in/Scripts/BS_ViewMasDirections.aspx'
MAKERS = ['a.rao@synthetic-finance.example', 'b.mehta@synthetic-finance.example', 'c.iyer@synthetic-finance.example',
          'd.khan@synthetic-finance.example', 'e.das@synthetic-finance.example', 'f.singh@synthetic-finance.example']
CHECKERS = ['cco@synthetic-finance.example', 'cfo@synthetic-finance.example', 'cro@synthetic-finance.example']

# (inst, dept, function, sub, freq, due_rule, event, short_title, question, risk, risk_type, comp_type, status, due_offset_days, evidence?)
ROWS = [
    ('GOV25', 'Governance & Administration', 'Board & Committees', 'Board meetings', 'Quarterly', 'At least once in each quarter', None, 'Board meeting frequency', 'Has the Board met at least once in the current quarter, with the minutes approved and signed?', 'High', 'Governance', 'Compliance', 'In Progress', 4, False),
    ('GOV25', 'Governance & Administration', 'Board & Committees', 'Committee constitution', 'Ongoing/Continuous', '', None, 'Risk Management Committee constitution', 'Is a Risk Management Committee of the Board constituted with the required composition?', 'Medium', 'Governance', 'Committee', 'Complied', None, True),
    ('GOV25', 'Governance & Administration', 'Fit & Proper', 'Fit and proper declarations', 'Annually', 'By 30 April each year', None, 'Annual fit and proper declarations', 'Have all directors submitted the annual fit and proper declaration and has the NRC reviewed them?', 'Medium', 'Governance', 'Compliance', 'Complied', None, True),
    ('GOV25', 'Governance & Administration', 'Policies & Codes', 'Policy review', 'Annually', 'Within 12 months of the last Board review', None, 'Annual review of Board-approved policies', 'Has the Board reviewed the policy register and approved the policies due for annual review?', 'Medium', 'Governance', 'Policies', 'Not Started', 21, False),
    ('GOV25', 'Governance & Administration', 'Fit & Proper', 'Regulator intimation', 'Event Based', 'Within 30 days of the change', 'Appointment of director / KMP / officer', 'Intimation of change in directors', 'Has the change in directors been intimated to RBI within the prescribed time?', 'High', 'Regulatory', 'Compliance', 'Not Started', None, False),
    ('KYC25', 'Compliance', 'KYC & AML', 'CDD', 'Ongoing/Continuous', '', None, 'Board-approved KYC policy', 'Does the NBFC have a Board-approved KYC policy covering customer acceptance, identification, monitoring and risk management?', 'High', 'Regulatory', 'Policies', 'Complied', None, True),
    ('KYC25', 'Compliance', 'KYC & AML', 'CDD', 'Half Yearly', 'Within 30 days of the end of each half-year', None, 'Periodic risk categorisation review', 'Has the risk categorisation of customers been reviewed for the half-year?', 'Medium', 'Regulatory; Fraud', 'Internal Compliance', 'In Progress', 36, True),
    ('KYC25', 'Compliance', 'Compliance Testing', 'Sample testing', 'Quarterly', 'Within 15 days from the end of each quarter', None, 'KYC sample testing', 'Has compliance tested a sample of new accounts for KYC completeness this quarter?', 'Medium', 'Regulatory', 'Internal Compliance', 'Returned for Rework', 6, True),
    ('CMPF26', 'Compliance', 'Regulatory Compliance', 'Compliance reporting', 'Annually', 'Within 3 months of the end of the financial year', None, 'Annual compliance plan', 'Has the Board approved the annual compliance plan, including the compliance testing programme?', 'High', 'Regulatory; Governance', 'Compliance', 'Complied', None, True),
    ('CMPF26', 'Compliance', 'Regulatory Liaison & Inspections', 'Reporting & Filing', 'Event Based', 'As specified in the inspection letter', 'Regulator direction, order or inspection report received', 'Inspection compliance report', 'Has the compliance report on the RBI inspection findings been submitted by the date specified?', 'High', 'Regulatory', 'Compliance', 'In Progress', -3, True),
    ('CMPF26', 'Compliance', 'Regulatory Compliance', 'Compliance reporting', 'Quarterly', 'Within 30 days from the end of each quarter', None, 'Quarterly compliance report to Board', 'Has the CCO placed the quarterly compliance report before the Board / Audit Committee?', 'Medium', 'Governance', 'Committee', 'Submitted for Review', 6, True),
    ('SR26', 'Returns & Reporting', 'Supervisory Returns', 'Monthly returns', 'Monthly', 'Within 7 days from the end of each month', None, 'Monthly return on important financial parameters', 'Has the monthly return been filed on CIMS within the due date?', 'High', 'Regulatory', 'Returns & Reporting', 'Not Started', 2, False),
    ('SR26', 'Returns & Reporting', 'Supervisory Returns', 'Quarterly returns', 'Quarterly', 'Within 30 days from the end of each quarter', None, 'Quarterly prudential return', 'Has the quarterly prudential return been validated and filed on CIMS?', 'High', 'Regulatory', 'Returns & Reporting', 'In Progress', 6, True),
    ('SR26', 'Returns & Reporting', 'Supervisory Returns', 'Return validation and sign-off', 'Quarterly', 'Before filing of each quarterly return', None, 'Return sign-off by authorised official', 'Has each quarterly return been signed off by the authorised official before filing?', 'Medium', 'Regulatory', 'Internal Compliance', 'Not Started', 6, False),
    ('SR26', 'Returns & Reporting', 'Supervisory Returns', 'Annual returns', 'Annually', 'Within 3 months from the end of the financial year', None, 'Annual statutory auditor certificate', 'Has the statutory auditor certificate been submitted for the year?', 'High', 'Regulatory', 'Returns & Reporting', 'Complied Late', None, True),
    ('SR26', 'Returns & Reporting', 'Supervisory Returns', 'Monthly returns', 'Monthly', 'Within 10 days from the end of each month', None, 'Monthly ALM-related return', 'Has the monthly liquidity return been filed within the due date?', 'Medium', 'Regulatory; Liquidity', 'Returns & Reporting', 'Not Complied', -15, False),
    ('CRM25', 'Credit', 'Credit Appraisal & Sanction', 'Policy & Framework', 'Ongoing/Continuous', '', None, 'Board-approved credit policy', 'Does the NBFC have a Board-approved credit policy with delegation of sanctioning powers?', 'High', 'Credit', 'Policies', 'Complied', None, True),
    ('CRM25', 'Credit', 'Credit Monitoring', 'SMA reporting', 'Weekly', 'Every Friday for the preceding week', None, 'Weekly SMA reporting to CICs', 'Has SMA status of borrowers been reported to the CICs for the week?', 'High', 'Credit; Regulatory', 'Returns & Reporting', 'In Progress', 1, True),
    ('CRM25', 'Credit', 'Credit Monitoring', 'Early warning signals', 'Monthly', 'Within 10 days from the end of each month', None, 'Early warning signal review', 'Has the early warning signal dashboard been reviewed and actions recorded for the month?', 'Medium', 'Credit', 'Internal Compliance', 'Not Started', 16, False),
    ('IRACP25', 'Credit', 'Asset Classification & Provisioning', 'NPA identification', 'Daily', 'Day-end classification run', None, 'Daily NPA flagging', 'Is the asset classification run executed daily with exceptions reviewed?', 'High', 'Credit', 'Compliance', 'Complied', None, True),
    ('IRACP25', 'Credit', 'Asset Classification & Provisioning', 'Provisioning', 'Quarterly', 'Before the quarterly financial results', None, 'Quarterly provisioning review', 'Have provisions been computed and reviewed by the Audit Committee for the quarter?', 'High', 'Credit', 'Committee', 'In Progress', 12, True),
    ('IRACP25', 'Risk Management', 'Credit Risk', 'Credit monitoring', 'Quarterly', 'Within 45 days from the end of each quarter', None, 'Upgradation review', 'Have upgraded accounts been checked for full arrears clearance before upgradation?', 'Medium', 'Credit', 'Internal Compliance', 'Not Started', 27, False),
    ('CRM25', 'Risk Management', 'Risk Governance', 'Risk appetite', 'Annually', 'Within 12 months of the last review', None, 'Risk appetite statement review', 'Has the Board reviewed the risk appetite statement for the year?', 'High', 'Governance; Credit', 'Policies', 'Not Started', 45, False),
    ('CRM25', 'Risk Management', 'Liquidity Risk', 'ALM', 'Monthly', 'Within 7 days from the end of each month', None, 'Liquidity risk tolerance monitoring', 'Has the ALCO reviewed liquidity gaps against the tolerance limits for the month?', 'Medium', 'Liquidity', 'Committee', 'Complied', None, True),
    ('FRM26', 'Risk Management', 'Operational Risk', 'Fraud risk', 'Event Based', 'Within 14 days of detection', 'Fraud detected', 'Fraud reporting to RBI', 'Has every fraud detected been reported to RBI within the prescribed time?', 'High', 'Fraud; Regulatory', 'Returns & Reporting', 'In Progress', 3, True),
    ('FRM26', 'Risk Management', 'Operational Risk', 'Fraud risk', 'Ongoing/Continuous', '', None, 'Fraud risk management policy', 'Is there a Board-approved fraud risk management policy with a Special Committee for monitoring?', 'High', 'Fraud', 'Policies', 'Complied', None, True),
    ('OUTS25', 'Operations & Management', 'Outsourcing', 'Materiality assessment', 'Annually', 'Within 12 months of the last assessment', None, 'Annual review of outsourcing arrangements', 'Has the annual review of material outsourcing arrangements been placed before the Board?', 'Medium', 'Third-party', 'Committee', 'In Progress', 19, True),
    ('OUTS25', 'Operations & Management', 'Outsourcing', 'Due diligence', 'Event Based', 'Before entering into the arrangement', 'Outsourcing or material contract entered into', 'Due diligence on service providers', 'Has due diligence been completed before onboarding each new service provider?', 'Medium', 'Third-party', 'Internal Compliance', 'Not Applicable', None, False),
    ('OUTS25', 'Operations & Management', 'Outsourcing', 'Ongoing monitoring', 'Half Yearly', 'Within 30 days of the end of each half-year', None, 'BCP testing of outsourced services', 'Have business continuity plans of material service providers been tested in the half-year?', 'Low', 'Operational', 'Internal Compliance', 'Not Started', 67, False),
    ('RBC25', 'Customer Service & Grievance', 'Grievance Redressal', 'Resolution and TAT', 'Monthly', 'Within 10 days from the end of each month', None, 'Grievance redressal MIS', 'Has the monthly grievance MIS, including complaints beyond TAT, been reviewed by senior management?', 'Medium', 'Conduct & Consumer', 'Internal Compliance', 'Submitted for Review', 9, True),
    ('RBC25', 'Customer Service & Grievance', 'Customer Communication & Disclosure', 'Process Execution', 'Ongoing/Continuous', '', None, 'Key Facts Statement', 'Is a Key Facts Statement issued to every borrower before execution of the loan agreement?', 'High', 'Conduct & Consumer', 'Compliance', 'Complied', None, True),
    ('RBC25', 'Customer Service & Grievance', 'Fair Practices', 'Monitoring & Review', 'Annually', 'Within 12 months of the last review', None, 'Fair Practices Code review', 'Has the Board reviewed the Fair Practices Code and its implementation for the year?', 'Medium', 'Conduct & Consumer', 'Policies', 'Not Started', 5, False),
    ('RBC25', 'Collections & Recovery', 'Collections Conduct', 'Process Execution', 'Ongoing/Continuous', '', None, 'Recovery agent conduct', 'Are recovery agents prohibited from calling borrowers outside the permitted hours, and is this monitored?', 'High', 'Conduct & Consumer', 'Compliance', 'Exempted', None, False),
    ('RBC25', 'Collections & Recovery', 'Recovery Agents Management', 'Records & Evidence', 'Quarterly', 'Within 15 days from the end of each quarter', None, 'Recovery agent empanelment review', 'Has the list of empanelled recovery agents been reviewed and published on the website this quarter?', 'Medium', 'Conduct & Consumer; Third-party', 'Internal Compliance', 'In Progress', -1, True),
    ('CYB26', 'Information Security', 'Incident Response', 'CERT-In / regulator reporting', 'Immediately', 'Within 6 hours of noticing the incident', 'Cyber or data security incident detected', 'Cyber incident reporting', 'Has every reportable cyber incident been reported to RBI and CERT-In within the time limit?', 'High', 'IT & Cyber; Regulatory', 'Returns & Reporting', 'Not Started', None, False),
    ('CYB26', 'Information Security', 'Cyber Security Governance', 'Cyber audit', 'Annually', 'Within 12 months of the last audit', None, 'Annual IS audit', 'Has the annual information systems audit been completed and placed before the IT Strategy Committee?', 'High', 'IT & Cyber', 'Compliance', 'In Progress', 8, True),
    ('CYB26', 'IT', 'IT Governance', 'IT policy', 'Annually', 'Within 12 months of the last review', None, 'IT policy review', 'Has the Board reviewed the IT and information security policies for the year?', 'Medium', 'IT & Cyber; Governance', 'Policies', 'Complied', None, True),
    ('CYB26', 'IT', 'IT Operations', 'Change management', 'Monthly', 'Within 5 days from the end of each month', None, 'Change management log review', 'Has the change advisory board log been reviewed for unauthorised changes this month?', 'Low', 'IT & Cyber', 'Internal Compliance', 'Not Started', 11, False),
    ('GOV25', 'Internal Audit', 'Audit Planning', 'Policy & Framework', 'Annually', 'Before the start of the financial year', None, 'Risk-based internal audit plan', 'Has the Audit Committee approved the risk-based internal audit plan for the year?', 'Medium', 'Governance', 'Committee', 'Complied', None, True),
    ('GOV25', 'Internal Audit', 'Audit Follow-up', 'Monitoring & Review', 'Quarterly', 'Within 30 days from the end of each quarter', None, 'Audit observation follow-up', 'Has the status of open audit observations been reported to the Audit Committee this quarter?', 'Medium', 'Governance', 'Committee', 'Waiting on Others', 6, True),
]
assert len(ROWS) == 40, len(ROWS)

CLOSED = {'Complied', 'Complied Late'}
GREYS = {'Not Applicable', 'Exempted', 'Withdrawn'}
SHORT_FREQ = {'Daily', 'Weekly', 'Fortnightly', 'Monthly'}


def rag_for(status, freq, due, risk, evidence, event_open, today=TODAY, period_days=None):
    """Default RAG rules — P04 FR-P04-230 (first match wins). Returns (rag, days_to_due)."""
    if status in GREYS:                                            # (1)
        return 'Grey', None
    if freq in ('Event Based',) and not event_open:                # (2) awaiting trigger
        return 'Blue', None
    if freq == 'Immediately' and not event_open:
        return 'Blue', None
    if status in CLOSED:                                           # (3)
        return 'Green', None
    if due is None:                                                # ongoing condition without a dated instance
        return ('Red', None) if status == 'Not Complied' else ('Green', None)
    days = (due - today).days
    if days < 0:                                                   # (4) overdue
        return 'Red', days
    if freq == 'Immediately':
        red_w, amb_w = 2, 3
    elif freq == 'Event Based':
        p = period_days or 30
        red_w, amb_w = max(2, math.ceil(0.25 * p)), math.ceil(0.5 * p)
    elif freq in SHORT_FREQ:
        red_w, amb_w = 3, 7
    else:
        red_w, amb_w = 7, 30
    if risk == 'High':                                             # High risk widens both windows by 50 %, rounded up
        red_w, amb_w = math.ceil(red_w * 1.5), math.ceil(amb_w * 1.5)
    if status == 'Returned for Rework' and days <= amb_w:          # (5)
        return 'Red', days
    if days <= red_w:                                              # (6)
        return 'Red', days
    if days <= amb_w or (not evidence and days <= amb_w):          # (7)
        return 'Amber', days
    return 'Green', days                                           # (8)


def build_records():
    seq = Counter()
    recs = []
    for i, r in enumerate(ROWS, 1):
        inst, dept, fn, sub, freq, rule, event, title, q, risk, rtype, ctype, status, off, evid = r
        seq[inst] += 1
        comp_id = f'RBI-NBFC-{inst}-{seq[inst] * 10:04d}'
        t, ref, issue, amend = INST[inst]
        due = TODAY + timedelta(days=off) if off is not None else None
        event_open = freq in ('Event Based', 'Immediately') and off is not None
        m = re.search(r'(\d+) days', rule or '')
        rag, days = rag_for(status, freq, due, risk, evid, event_open, period_days=int(m.group(1)) if m else None)
        maker = MAKERS[i % len(MAKERS)]
        upd = amend if (amend and i % 3 == 0) else issue
        act = None
        if status == 'Complied':
            act = TODAY - timedelta(days=10 + i)
        if status == 'Complied Late':
            act = TODAY - timedelta(days=40)
        rec = OrderedDict(
            sno=i, comp_id=comp_id, question=q, clause_ref=f'Para {4 + (i * 7) % 60} (illustrative)', department=dept, frequency=freq,
            act_date=act, status=status, risk_level=risk, authority='RBI', instrument_name=t, instrument_ref=ref, instrument_issue_date=issue,
            instrument_amended_date=amend or issue, question_updated_date=upd, change_status='Amended' if upd != issue else 'New',
            change_note=(f'Amended vide notification dated {upd:%d-%b-%Y}; previous wording retained in History (synthetic)' if upd != issue else None),
            source_url=URL, applic_layer='All Layers' if i % 4 else 'Middle Layer (NBFC-ML); Upper Layer (NBFC-UL)',
            applic_type='All NBFCs (ALL)', conditions='Applies to all NBFCs. (SYNTHETIC sample row)', applic_from=None, applic_to=None,
            ext_customer_interface=('With customer interface' if i % 3 == 0 else None), ext_public_funds=None, ext_asset_size=None,
            short_title=title, explanation='SYNTHETIC sample text: plain-language explanation of the actionable would appear here.',
            dept_function=fn, dept_subfunction=sub, compliance_type=ctype, due_rule=rule or None, event_trigger=event,
            filing_details=('RBI - CIMS (Supervisory Returns)' if ctype == 'Returns & Reporting' else 'Board / Management'),
            filing_form=None, related_docs=None, weblinks=None, consequences='SYNTHETIC: supervisory action and monetary penalty under the RBI Act, 1934.',
            penalty_ref='s.58G, RBI Act, 1934', risk_type=rtype, risk_ids=('RSK-SYNF-000%02d' % (i % 9 + 1)) if i % 2 else None,
            gov_ids=('GOV-SYNF-MTG-000%02d' % (i % 7 + 1)) if dept == 'Governance & Administration' else None,
            evidence='SYNTHETIC: board minutes / filing acknowledgement / register extract', control_activity='SYNTHETIC: maker steps.',
            checker_points='SYNTHETIC: checker points.', audit_test=None, tags=None,
            maker=maker, maker_status=('Evidence Uploaded' if evid else ('Not Started' if status == 'Not Started' else 'In Progress')),
            evidence_ref=(f'VAULT/{comp_id}/2026-Q3/01.pdf' if evid else None), evidence_date=(TODAY - timedelta(days=5) if evid else None),
            maker_comments=('Awaiting data from Finance (synthetic)' if status == 'Waiting on Others' else None),
            checker=CHECKERS[i % 3], checker_decision=('Approve' if status in CLOSED else ('Reject — Return to Maker' if status == 'Returned for Rework' else None)),
            approval_date=(act if status in CLOSED else None), checker_comments=('Evidence incomplete — resubmit (synthetic)' if status == 'Returned for Rework' else None),
            auditor='ia@synthetic-finance.example' if i % 5 == 0 else None, audit_result=('Satisfactory' if i % 10 == 0 else ('Not Tested' if i % 5 == 0 else None)),
            auditor_comments=None, accountable_owner=f'Head — {dept}', next_due=due, days_to_due=days, rag=rag, question_version=1 + (upd != issue),
            overlay='No')
        recs.append(rec)
    return recs


# ------------------------------------------------------------------ Excel export
THIN = Side(style='thin', color='FF9E9E9E')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
EXPORT_COLS = K.EXPORT_DEFAULT_KEYS  # list view + next due + days + RAG + maker
HDR = {c['key']: c['header'] for c in K.family_columns('RBI-NBFC')}


def fill(h):
    return PatternFill('solid', fgColor='FF' + h)


def put_list_sheet(ws, recs, group_key=None, title=None):
    """Write the default export columns; if group_key, insert group headers + subtotal rows with row outline grouping."""
    r = 1
    ws.cell(row=r, column=1, value=title or 'Compliance list').font = Font(bold=True, size=13)
    ws.cell(row=r, column=6, value='SYNTHETIC SAMPLE — not legal content').font = Font(bold=True, color='FFC00000')
    r = 2
    cols = EXPORT_COLS
    widths = {'sno': 6, 'comp_id': 22, 'question': 58, 'clause_ref': 16, 'department': 22, 'frequency': 16, 'act_date': 12, 'status': 18,
              'risk_level': 9, 'next_due': 12, 'days_to_due': 8, 'rag': 8, 'maker': 30}
    for j, k in enumerate(cols, 1):
        c = ws.cell(row=r, column=j, value=HDR[k])
        c.font = Font(bold=True, size=10)
        c.fill = fill('FFF2CC' if k in K.LIST_VIEW_KEYS else ('EDEDED' if k != 'maker' else 'E4DFEC'))
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BOX
        ws.column_dimensions[get_column_letter(j)].width = widths[k]
    r = 3
    rag_col = get_column_letter(cols.index('rag') + 1)
    groups = OrderedDict()
    for rec in recs:
        groups.setdefault(rec[group_key] if group_key else '', []).append(rec)
    for g, items in groups.items():
        if group_key:
            ws.cell(row=r, column=1, value=f'{HDR[group_key]}: {g}').font = Font(bold=True, size=11)
            for j in range(1, len(cols) + 1):
                ws.cell(row=r, column=j).fill = fill('D9E1F2')
            r += 1
        first = r
        for rec in items:
            for j, k in enumerate(cols, 1):
                v = rec[k]
                c = ws.cell(row=r, column=j, value=v)
                c.alignment = WRAP
                c.border = BOX
                c.font = Font(size=9)
                if isinstance(v, date):
                    c.number_format = 'DD-MMM-YYYY'
            if group_key:
                ws.row_dimensions[r].outline_level = 1
            r += 1
        if group_key:
            cnt = Counter(x['rag'] for x in items)
            ws.cell(row=r, column=1, value=f'Subtotal — {g}: {len(items)} items · ' + ' · '.join(f'{k} {cnt.get(k, 0)}' for k in K.RAG)).font = Font(bold=True, italic=True, size=9)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
            r += 2
    for nm, hx in K.RAG_FILL.items():
        ws.conditional_formatting.add(f'{rag_col}3:{rag_col}{r}', FormulaRule(formula=[f'${rag_col}3="{nm}"'], fill=fill(hx)))
        dl = get_column_letter(cols.index('days_to_due') + 1)
        ws.conditional_formatting.add(f'{dl}3:{dl}{r}', FormulaRule(formula=[f'${rag_col}3="{nm}"'], fill=fill(hx)))
    ws.freeze_panes = 'D3'
    if not group_key:
        ws.auto_filter.ref = f'A2:{get_column_letter(len(cols))}{r - 1}'
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '2:2'
    ws.oddFooter.left.text = EXPORT_ID
    ws.oddFooter.center.text = 'SYNTHETIC SAMPLE'
    ws.oddFooter.right.text = 'Page &P of &N'


def build_xlsx(recs):
    wb = openpyxl.load_workbook(TEMPLATE)
    reg = wb[K.SHEET]
    keys = [reg.cell(row=2, column=j).value for j in range(1, reg.max_column + 1)]
    info_keys = [c['key'] for c in K.family_columns('RBI-NBFC')]
    assert keys == [HDR[k] for k in info_keys]
    for i, rec in enumerate(recs, 3):
        for j, k in enumerate(info_keys, 1):
            v = rec.get(k)
            c = reg.cell(row=i, column=j, value=v)
            if isinstance(v, date):
                c.number_format = 'DD-MMM-YYYY'
    reg.auto_filter.ref = f'A2:{get_column_letter(len(info_keys))}{len(recs) + 2}'
    reg._print_area = _WholeColumns(f"'{K.SHEET}'!$A:$I")
    reg.print_title_rows = '1:2'

    # Summary
    sm = wb.create_sheet('Summary', 0)
    sm['A1'] = 'Working Register Export — Summary'
    sm['A1'].font = Font(bold=True, size=14)
    sm['A2'] = 'SYNTHETIC SAMPLE — illustrative test data only; not legal content and not a real entity'
    sm['A2'].font = Font(bold=True, color='FFC00000')
    meta = [('Tenant', TENANT), ('Template', 'LDC-TPL-RBI-NBFC-v2.0'), ('Export ID (ledger event)', EXPORT_ID), ('Export date', TODAY),
            ('Filter', 'Entity = Synthetic Finance Ltd; Period = 2026-Q3 / Oct-2026 instances; all statuses'),
            ('Group by', 'Department (one sheet per Department follows)'), ('Columns', 'Default export (D-50): list view + Next Due Date + Days to Due + RAG + Maker'),
            ('Items', len(recs)), ('RAG rules', 'Platform defaults (Appendix E §A.12); High risk widens windows by 50%')]
    for i, (k, v) in enumerate(meta, 4):
        sm.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = sm.cell(row=i, column=2, value=v)
        c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        sm.merge_cells(start_row=i, start_column=2, end_row=i, end_column=10)
        if isinstance(v, date):
            c.number_format = 'DD-MMM-YYYY'
    r = 4 + len(meta) + 1

    def matrix(r, key, label):
        sm.cell(row=r, column=1, value=f'{label} × RAG').font = Font(bold=True, size=12)
        r += 1
        heads = [label] + K.RAG + ['Total']
        for j, h in enumerate(heads, 1):
            c = sm.cell(row=r, column=j, value=h)
            c.font = Font(bold=True)
            c.fill = fill(K.RAG_FILL.get(h, 'D9E1F2'))
            c.border = BOX
        r += 1
        vals = OrderedDict()
        for rec in recs:
            vals.setdefault(rec[key], Counter())[rec['rag']] += 1
        start = r
        for g, cnt in sorted(vals.items(), key=lambda kv: (kv[0].startswith('No '), kv[0])):
            row = [g] + [cnt.get(k, 0) for k in K.RAG]
            for j, v in enumerate(row, 1):
                sm.cell(row=r, column=j, value=v).border = BOX
            sm.cell(row=r, column=len(heads), value=f'=SUM(B{r}:F{r})').border = BOX
            r += 1
        sm.cell(row=r, column=1, value='Total').font = Font(bold=True)
        for j in range(2, len(heads) + 1):
            L = get_column_letter(j)
            c = sm.cell(row=r, column=j, value=f'=SUM({L}{start}:{L}{r - 1})')
            c.font = Font(bold=True)
            c.border = BOX
        return r + 2

    r = matrix(r, 'department', 'Department')
    for rec in recs:
        rec['due_month'] = rec['next_due'].strftime('%Y-%m') if rec['next_due'] else 'No fixed due date'
        rec['primary_risk_type'] = rec['risk_type'].split('; ')[0]
    r = matrix(r, 'primary_risk_type', 'Risk Type (primary)')
    r = matrix(r, 'due_month', 'Due Month')
    sm.column_dimensions['A'].width = 34
    for L in 'BCDEFGHIJ':
        sm.column_dimensions[L].width = 9
    sm.page_setup.orientation = 'landscape'
    sm.page_setup.paperSize = sm.PAPERSIZE_A4
    sm.page_setup.fitToWidth, sm.page_setup.fitToHeight = 1, 0
    sm.sheet_properties.pageSetUpPr.fitToPage = True

    # grouped list (all departments) + one sheet per department
    recs_sorted = sorted(recs, key=lambda x: (x['department'], x['next_due'] or date(2100, 1, 1)))
    gl = wb.create_sheet('Compliance List (grouped)', 1)
    put_list_sheet(gl, recs_sorted, 'department', 'Compliance list — grouped by Department (default export columns)')
    pos = 3
    for dept in OrderedDict.fromkeys(x['department'] for x in recs_sorted):
        ws = wb.create_sheet(dept[:31].replace('/', '-'), pos)
        put_list_sheet(ws, [x for x in recs_sorted if x['department'] == dept], None, f'Department: {dept}')
        pos += 1
    ti = wb['_Template_Info']  # export identity block (Appendix E §A.11.3); appended after the contract and legacy map
    r0 = ti.max_row + 2
    ti.cell(row=r0, column=1, value='EXPORT IDENTITY').font = Font(bold=True)
    for i, (k, v) in enumerate([('export_id', EXPORT_ID), ('tenant_id_hash', 'sha256:5f1c…(synthetic)'), ('draft', 'false'),
                                ('export_kind', 'WORKING_REGISTER'), ('exported_at', '2026-10-25T18:30:00+05:30')], r0 + 1):
        ti.cell(row=i, column=1, value=k).font = Font(bold=True)
        ti.cell(row=i, column=4, value=v)
    ex = wb.create_sheet('_Export_Info')
    for i, (k, v) in enumerate([('export_id', EXPORT_ID), ('template_code', 'LDC-TPL-RBI-NBFC-v2.0'), ('contract_version', '2.0'),
                                ('export_kind', 'WORKING_REGISTER'), ('draft', 'false'), ('tenant_id_hash', 'sha256:5f1c…(synthetic)'),
                                ('filter', 'entity = Synthetic Finance Ltd; instances 2026-Q3 / 2026-10; status = all'), ('group_by', 'department'),
                                ('exported_at', '2026-10-25T18:30:00+05:30'), ('exported_by', 'cco@synthetic-finance.example'),
                                ('row_count', len(recs)), ('comp_ids', '; '.join(x['comp_id'] for x in recs)),
                                ('ledger', 'One export event per CompID written to the ledger (NN-57)'),
                                ('re_import', 'Re-import updates the same CompIDs; User/System groups are ignored on content import'),
                                ('synthetic', 'YES — test data only')], 1):
        ex.cell(row=i, column=1, value=k).font = Font(bold=True)
        ex.cell(row=i, column=2, value=v).alignment = WRAP
    ex.column_dimensions['A'].width = 18
    ex.column_dimensions['B'].width = 120
    ex.protection.sheet = True
    for rec in recs:
        rec.pop('due_month', None)
        rec.pop('primary_risk_type', None)
    wb.active = 0
    path = os.path.join(OUT, 'SAMPLE - Working Register Export - RBI-NBFC.xlsx')
    wb.save(path)
    return path


# ------------------------------------------------------------------ PDF export
def build_pdf(recs):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether

    pdfmetrics.registerFont(TTFont('DJ', '/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf'))
    pdfmetrics.registerFont(TTFont('DJB', '/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed-Bold.ttf'))
    st = ParagraphStyle('b', fontName='DJ', fontSize=7.4, leading=9)
    stb = ParagraphStyle('bb', fontName='DJB', fontSize=7.4, leading=9)
    h1 = ParagraphStyle('h1', fontName='DJB', fontSize=17, leading=21)
    h2 = ParagraphStyle('h2', fontName='DJB', fontSize=11.5, leading=15, spaceBefore=4, spaceAfter=3)
    small = ParagraphStyle('s', fontName='DJ', fontSize=8.5, leading=11)
    red = ParagraphStyle('r', fontName='DJB', fontSize=9, leading=12, textColor=colors.HexColor('#C00000'))
    path = os.path.join(OUT, 'SAMPLE - Compliance List Export - RBI-NBFC.pdf')

    def footer(cv, doc):
        cv.saveState()
        cv.setFont('DJ', 7)
        cv.drawString(12 * mm, 7 * mm, f'{EXPORT_ID} · LDC-TPL-RBI-NBFC-v2.0 · exported 25-Oct-2026 18:30 IST')
        cv.setFillColor(colors.HexColor('#C00000'))
        cv.drawCentredString(landscape(A4)[0] / 2 + 20 * mm, 7 * mm, 'SYNTHETIC SAMPLE — not legal content')
        cv.setFillColor(colors.black)
        cv.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f'Page {doc.page}')
        cv.restoreState()

    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=13 * mm,
                            title='Compliance List Export — RBI-NBFC (SYNTHETIC SAMPLE)', author='LawDocsComply')
    story = [Paragraph('Compliance List Export — grouped by Department', h1), Spacer(1, 3 * mm),
             Paragraph('SYNTHETIC SAMPLE — illustrative test data generated for the LawDocsComply build pack v1.2; not legal content and not a real entity.', red),
             Spacer(1, 3 * mm)]
    meta = [['Tenant', TENANT], ['Export ID', EXPORT_ID + ' (one ledger event per CompID)'], ['Export date', '25-Oct-2026'],
            ['Filter', 'Entity = Synthetic Finance Ltd; current instances (2026-Q3 / Oct-2026); all statuses'], ['Group by', 'Department'],
            ['Columns', 'List view (S. No., CompID, Question, Clause, Department, Due Date & Frequency, Date of Compliance, Status, Risk) + Next Due, Days to Due, RAG, Maker'],
            ['RAG rules', 'Platform defaults (P04 FR-P04-230, first match wins): Grey = not applicable / exempted / withdrawn; Blue = event-based awaiting trigger; Green = complied; '
                          'Red = overdue, returned within the Amber window, or inside the Red window (Immediately 2 days; daily–monthly 3; quarterly and longer 7; event-based 25% of the period left); '
                          'Amber = inside the Amber window (3 / 7 / 30 days; event-based 50% left); High risk widens both windows by 50%.']]
    t = Table([[Paragraph(a, stb), Paragraph(b, st)] for a, b in meta], colWidths=[35 * mm, 230 * mm])
    t.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.3, colors.grey), ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#D9E1F2')),
                           ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
    story += [t, Spacer(1, 5 * mm), Paragraph('RAG summary by Department', h2)]
    depts = OrderedDict()
    for rec in sorted(recs, key=lambda x: x['department']):
        depts.setdefault(rec['department'], []).append(rec)
    rows = [['Department'] + K.RAG + ['Total']]
    tot = Counter()
    for d, items in depts.items():
        cnt = Counter(x['rag'] for x in items)
        tot.update(cnt)
        rows.append([d] + [cnt.get(k, 0) for k in K.RAG] + [len(items)])
    rows.append(['Total'] + [tot.get(k, 0) for k in K.RAG] + [len(recs)])
    ts = Table(rows, colWidths=[70 * mm] + [18 * mm] * 6)
    sty = [('GRID', (0, 0), (-1, -1), 0.3, colors.grey), ('FONTNAME', (0, 0), (-1, -1), 'DJ'), ('FONTSIZE', (0, 0), (-1, -1), 8.5),
           ('FONTNAME', (0, 0), (-1, 0), 'DJB'), ('FONTNAME', (0, -1), (-1, -1), 'DJB'), ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
           ('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#D9E1F2')), ('BACKGROUND', (-1, 0), (-1, 0), colors.HexColor('#D9E1F2'))]
    for j, k in enumerate(K.RAG, 1):
        sty.append(('BACKGROUND', (j, 0), (j, 0), colors.HexColor('#' + K.RAG_FILL[k])))
    ts.setStyle(TableStyle(sty))
    story += [ts, PageBreak()]

    heads = ['S. No.', 'CompID', 'Compliance Question / Actionable', 'Clause / Para', 'Due Date & Frequency', 'Date of Compliance', 'Status', 'Risk',
             'Next Due', 'Days', 'RAG', 'Maker']
    widths = [9, 34, 59, 20, 32, 19, 21, 14, 20, 11, 13, 0]
    widths[-1] = 273 - sum(widths)
    fmt = lambda d: d.strftime('%d-%b-%Y') if d else '—'  # noqa: E731
    for d, items in depts.items():
        items = sorted(items, key=lambda x: x['next_due'] or date(2100, 1, 1))
        data = [[Paragraph(h, stb) for h in heads]]
        for x in items:
            freq = x['frequency'] + (f' — {x["due_rule"]}' if x['due_rule'] else '')
            data.append([Paragraph(str(x['sno']), st), Paragraph(x['comp_id'], st), Paragraph(x['question'], st), Paragraph(x['clause_ref'], st),
                         Paragraph(freq, st), Paragraph(fmt(x['act_date']), st), Paragraph(x['status'], st), Paragraph(x['risk_level'], st),
                         Paragraph(fmt(x['next_due']), st), Paragraph('' if x['days_to_due'] is None else str(x['days_to_due']), st),
                         Paragraph(x['rag'], stb), Paragraph(x['maker'].split('@')[0] + '@…', st)])
        cnt = Counter(x['rag'] for x in items)
        data.append([Paragraph(f'Subtotal — {d}: {len(items)} items · ' + ' · '.join(f'{k} {cnt.get(k, 0)}' for k in K.RAG), stb)] + [''] * 11)
        tb = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        sty = [('GRID', (0, 0), (-1, -2), 0.3, colors.grey), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#FFF2CC')),
               ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('SPAN', (0, -1), (-1, -1)), ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#EDEDED')),
               ('BOX', (0, -1), (-1, -1), 0.3, colors.grey),
               ('LEFTPADDING', (0, 0), (-1, -1), 2.5), ('RIGHTPADDING', (0, 0), (-1, -1), 2)]
        for i, x in enumerate(items, 1):
            sty.append(('BACKGROUND', (10, i), (10, i), colors.HexColor('#' + K.RAG_FILL[x['rag']])))
        tb.setStyle(TableStyle(sty))
        story += [KeepTogether([Paragraph(f'Department: {d}', h2), tb]), Spacer(1, 4 * mm)]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def main():
    os.makedirs(OUT, exist_ok=True)
    recs = build_records()
    print(Counter(r['rag'] for r in recs))
    print(build_xlsx([dict(r) for r in recs]))
    print(build_pdf(recs))


if __name__ == '__main__':
    main()
